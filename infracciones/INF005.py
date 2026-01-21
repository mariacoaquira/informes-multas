# --- BIBLIOTECAS ---
import streamlit as st
import pandas as pd
import io
from babel.dates import format_date
from num2words import num2words
from docx import Document
from docxcompose.composer import Composer
from docxtpl import DocxTemplate, RichText
from datetime import date, timedelta
import holidays # Aún puede ser útil para calcular fechas base

# --- IMPORTACIONES DE MÓDULOS PROPIOS ---
from textos_manager import obtener_fuente_formateada
from funciones import (create_table_subdoc, create_main_table_subdoc, texto_con_numero,
                     create_footnotes_subdoc, create_consolidated_bi_table_subdoc,
                     create_personal_table_subdoc, redondeo_excel, format_decimal_dinamico) # <-- AÑADIR ESTAS DOS
from sheets import (calcular_beneficio_ilicito, calcular_multa, 
                    descargar_archivo_drive, 
                    calcular_beneficio_ilicito_extemporaneo)
# --- IMPORTAR EL MÓDULO DE CAPACITACIÓN ---
from modulos.calculo_capacitacion import calcular_costo_capacitacion
from funciones import create_main_table_subdoc, create_table_subdoc, texto_con_numero, create_footnotes_subdoc, format_decimal_dinamico, redondeo_excel

# ---------------------------------------------------------------------
# FUNCIÓN AUXILIAR: CÁLCULO CE COMPLETO (CE1 + CE2) POR EXTREMO
# (Versión final con lógica CE1 interna y fecha incumplimiento única - CORREGIDA INDENTACIÓN)
# ---------------------------------------------------------------------

def _calcular_costo_evitado_extremo_inf005(datos_comunes, datos_hecho_general, extremo_data):
    """
    Calcula el CE completo (CE1 + CE2 condicional) para un único extremo de INF005.
    - USA SIEMPRE fecha_incumplimiento para calcular los montos de CE1 y CE2.
    - SUMA CE2 al costo para BI ('ce_soles_para_bi') SÓLO si tipo_extremo es 'No remitió'.
    Devuelve un diccionario con todos los componentes y totales.
    """
    result = {
        'ce1_data_raw': [], 'ce1_soles': 0.0, 'ce1_dolares': 0.0,
        'ce2_data_raw': [], 'ce2_soles_calculado': 0.0, 'ce2_dolares_calculado': 0.0,
        'ce_soles_para_bi': 0.0, 'ce_dolares_para_bi': 0.0,
        'aplicar_ce2_a_bi': False,
        'ids_anexos': set(),
        'fuentes': {'ce1': {}, 'ce2': {}}, # Para placeholders de fuentes CE1 y CE2
        'error': None
    }
    try:
        # --- Datos del Extremo y Generales ---
        tipo_incumplimiento = extremo_data.get('tipo_extremo')
        fecha_incumplimiento_extremo = extremo_data.get('fecha_incumplimiento')
        num_personal_ce2 = datos_hecho_general.get('num_personal_capacitacion', 0) # Tomar de datos generales

        # Validar fecha clave para el cálculo
        if not fecha_incumplimiento_extremo:
            raise ValueError("Falta la fecha de incumplimiento del extremo.")

        # --- INICIO CORRECCIÓN 1: Unificar Fecha y Fuentes ---
        df_ind = datos_comunes.get('df_indices')
        if df_ind is None: raise ValueError("Faltan df_indices en datos_comunes.")
        
        fecha_final_dt = pd.to_datetime(fecha_incumplimiento_extremo, errors='coerce')
        if pd.isna(fecha_final_dt): raise ValueError(f"Fecha de incumplimiento inválida: {fecha_incumplimiento_extremo}")

        ipc_row = df_ind[df_ind['Indice_Mes'].dt.to_period('M') == fecha_final_dt.to_period('M')]
        if ipc_row.empty:
            raise ValueError(f"No se encontró IPC/TC para la fecha de incumplimiento {fecha_final_dt.strftime('%B %Y')}")
            
        ipc_inc, tc_inc = ipc_row.iloc[0]['IPC_Mensual'], ipc_row.iloc[0]['TC_Mensual']
        if pd.isna(ipc_inc) or pd.isna(tc_inc) or tc_inc == 0:
            raise ValueError("Valores IPC/TC inválidos para la fecha de incumplimiento.")

        # Guardar las fuentes unificadas en el resultado principal
        result['fuentes']['fi_mes'] = format_date(fecha_final_dt, "MMMM 'de' yyyy", locale='es')
        result['fuentes']['fi_ipc'] = float(ipc_inc)
        result['fuentes']['fi_tc'] = float(tc_inc)
        # --- FIN CORRECCIÓN 1 ---

        # --- a. Calcular CE1 (Remisión - Lógica Interna con Horas Fijas) ---
        # *** FECHA DE CÁLCULO CE1: Siempre fecha_incumplimiento_extremo ***
        fecha_calculo_ce1 = fecha_final_dt # Usar el datetime object

        # --- INICIO Lógica interna para calcular CE1 ---
        # (Esta función interna está indentada correctamente bajo el 'try' principal)
        def _calcular_ce1_interno(datos_comunes_ce1, fecha_final):
            res_int = {'items_calculados': [], 'error': None, 'fuentes': {}}
            try:
                # (Desempaquetado y validaciones - indentación correcta)
                df_items_inf = datos_comunes_ce1.get('df_items_infracciones')
                df_costos = datos_comunes_ce1.get('df_costos_items')
                df_coti = datos_comunes_ce1.get('df_coti_general')
                df_sal = datos_comunes_ce1.get('df_salarios_general')
                df_ind = datos_comunes_ce1.get('df_indices')
                id_rubro_ce1 = datos_comunes_ce1.get('id_rubro_seleccionado')
                id_inf_ce1 = 'INF005'
                if any(df is None for df in [df_items_inf, df_costos, df_coti, df_sal, df_ind]): raise ValueError("Faltan DataFrames CE1.")

                fecha_final_dt_ce1 = pd.to_datetime(fecha_final, errors='coerce')
                if pd.isna(fecha_final_dt_ce1): raise ValueError(f"Fecha inválida CE1: {fecha_final}")
                ipc_row_ce1 = df_ind[df_ind['Indice_Mes'].dt.to_period('M') == fecha_final_dt_ce1.to_period('M')]
                if ipc_row_ce1.empty: raise ValueError(f"Sin IPC/TC CE1 para {fecha_final_dt_ce1.strftime('%B %Y')}")
                ipc_inc_ce1, tc_inc_ce1 = ipc_row_ce1.iloc[0]['IPC_Mensual'], ipc_row_ce1.iloc[0]['TC_Mensual']
                if pd.isna(ipc_inc_ce1) or pd.isna(tc_inc_ce1) or tc_inc_ce1 == 0: raise ValueError("IPC/TC inválidos CE1")

                # (Inicializaciones - indentación correcta)
                fuentes_ce1 = {'placeholders_dinamicos': {}}; items_ce1 = []; sal_capturado = False
                receta_ce1 = df_items_inf[df_items_inf['ID_Infraccion'] == id_inf_ce1]
                if receta_ce1.empty: raise ValueError(f"No hay receta CE1 para {id_inf_ce1}")

                # (Bucle for - indentación correcta)
                for _, item_receta in receta_ce1.iterrows():
                    # (Condición Tipo_Costo - indentación correcta)
                    if item_receta.get('Tipo_Costo') != 'Remision':
                        continue

                    # (Lógica de búsqueda de costo - indentación correcta)
                    id_item = item_receta['ID_Item_Infraccion']; desc_item = item_receta.get('Nombre_Item', 'N/A')
                    costos_posibles = df_costos[df_costos['ID_Item_Infraccion'] == id_item].copy();
                    if costos_posibles.empty: continue
                    tipo_item = item_receta.get('Tipo_Item'); df_candidatos = pd.DataFrame()
                    if tipo_item == 'Variable':
                        id_rubro_str = str(id_rubro_ce1) if id_rubro_ce1 else ''; df_candidatos = costos_posibles[costos_posibles['ID_Rubro'].astype(str).str.contains(fr'\b{id_rubro_str}\b', regex=True, na=False)].copy() if id_rubro_str else pd.DataFrame()
                        if df_candidatos.empty: df_candidatos = costos_posibles[costos_posibles['ID_Rubro'].isin(['', 'nan', None])].copy()
                    elif tipo_item == 'Fijo': df_candidatos = costos_posibles.copy()
                    if df_candidatos.empty: continue
                    fechas_fuente = []
                    for _, cand in df_candidatos.iterrows(): # Bucle interno
                        id_gen = cand['ID_General']; fecha_f = pd.NaT
                        if pd.notna(id_gen):
                            if 'SAL' in id_gen: f = df_sal[df_sal['ID_Salario'] == id_gen]; fecha_f = pd.to_datetime(f"{int(f.iloc[0]['Costeo_Salario'])}-12-31", errors='coerce') if not f.empty else pd.NaT
                            elif 'COT' in id_gen: f = df_coti[df_coti['ID_Cotizacion'] == id_gen]; fecha_f = pd.to_datetime(f.iloc[0]['Fecha_Costeo'], errors='coerce') if not f.empty else pd.NaT
                        fechas_fuente.append(fecha_f)
                    df_candidatos['Fecha_Fuente'] = fechas_fuente; df_candidatos.dropna(subset=['Fecha_Fuente'], inplace=True)
                    if df_candidatos.empty: continue
                    fecha_naive = fecha_final_dt_ce1.tz_localize(None) if fecha_final_dt_ce1.tzinfo else fecha_final_dt_ce1
                    df_candidatos['Fecha_Fuente_Naive'] = df_candidatos['Fecha_Fuente'].apply(lambda x: x.tz_localize(None) if pd.notna(x) and x.tzinfo else x)
                    df_candidatos['Diferencia_Dias'] = (df_candidatos['Fecha_Fuente_Naive'] - fecha_naive).dt.days.abs()
                    costo_final = df_candidatos.loc[df_candidatos['Diferencia_Dias'].idxmin()]

                    # (IPC/TC Costeo - indentación correcta)
                    id_gen = costo_final['ID_General']; fecha_f = costo_final['Fecha_Fuente']
                    ipc_cost, tc_cost = 0.0, 0.0
                    if pd.notna(id_gen) and 'SAL' in id_gen:
                        idx_anio = df_ind[df_ind['Indice_Mes'].dt.year == fecha_f.year]
                        ipc_cost, tc_cost = (float(idx_anio['IPC_Mensual'].mean()), float(idx_anio['TC_Mensual'].mean())) if not idx_anio.empty else (0.0, 0.0)
                        
                        # --- NUEVO: Capturar Fuente Salario e IPC Promedio ---
                        f_row = df_sal[df_sal['ID_Salario'] == id_gen]
                        if not f_row.empty:
                            fuentes_ce1['fuente_salario'] = f_row.iloc[0].get('Fuente_Salario', '')
                            fuentes_ce1['pdf_salario'] = f_row.iloc[0].get('PDF_Salario', '')
                            # Placeholder solicitado para el informe
                            fuentes_ce1['texto_ipc_costeo_salario'] = f"Promedio {fecha_f.year}, IPC = {ipc_cost:,.6f}"
                            sal_capturado = True
                    elif pd.notna(id_gen) and 'COT' in id_gen: idx_row = df_ind[df_ind['Indice_Mes'].dt.to_period('M') == fecha_f.to_period('M')]; ipc_cost, tc_cost = (float(idx_row.iloc[0]['IPC_Mensual']), float(idx_row.iloc[0]['TC_Mensual'])) if not idx_row.empty else (0.0, 0.0)
                    if ipc_cost == 0 or pd.isna(ipc_cost): continue

                    # (Capturar Fuentes - indentación correcta y estilo mejorado)
                    if pd.notna(id_gen):
                        # Fuente de Cotización
                        if 'COT' in id_gen:
                            f_row = df_coti[df_coti['ID_Cotizacion'] == id_gen]
                            if not f_row.empty:
                                sust = f_row.iloc[0].get('Fuente_Cotizacion')
                                if sust:
                                    # setdefault crea la lista si no existe, luego append añade
                                    fuentes_ce1.setdefault('fuente_coti', []).append(sust)
                        # Fuente de Salario (solo la primera vez)
                        elif 'SAL' in id_gen and not sal_capturado:
                            f_row = df_sal[df_sal['ID_Salario'] == id_gen]
                            if not f_row.empty:
                                fuentes_ce1['fuente_salario'] = f_row.iloc[0].get('Fuente_Salario', '')
                                fuentes_ce1['pdf_salario'] = f_row.iloc[0].get('PDF_Salario', '') # Asegúrate que 'PDF_Salario' es el nombre correcto
                                sal_capturado = True # Marcar que ya se capturó

                    # Sustento Profesional (fuera del if pd.notna(id_gen))
                    if "Profesional" in desc_item:
                        fuentes_ce1['sustento_item_profesional'] = costo_final.get('Sustento_Item', '')

                    # Placeholders dinámicos (fuera del if pd.notna(id_gen))
                    try:
                        # Corregido: Se eliminó la comilla extra (') al final del f-string
                        key_ph = f"fuente_{desc_item.split()[0].lower().replace(':','')}"
                        fecha_ph = format_date(fecha_f, 'MMMM yyyy', locale='es').lower()
                        texto_ph = f"{desc_item}:\n{fecha_ph}, IPC={ipc_cost:,.3f}"
                        # Crear el diccionario si no existe antes de asignar
                        if 'placeholders_dinamicos' not in fuentes_ce1:
                            fuentes_ce1['placeholders_dinamicos'] = {}
                        fuentes_ce1['placeholders_dinamicos'][key_ph] = texto_ph
                    except Exception as e:
                        # Es buena idea al menos imprimir el error si ocurre
                        print(f"Advertencia: No se pudo crear placeholder dinámico para '{desc_item}': {e}")
                        pass # Continuar si falla la creación del placeholder

                    # (Calcular Montos Horas Fijas - indentación correcta)
                    costo_orig=float(costo_final.get('Costo_Unitario_Item', 0.0)); moneda=costo_final.get('Moneda_Item')
                    if moneda!='S/' and (tc_cost==0 or pd.isna(tc_cost)): continue
                    precio_s=costo_orig if moneda=='S/' else costo_orig*tc_cost
                    factor=redondeo_excel(ipc_inc_ce1/ipc_cost, 3) if ipc_cost>0 else 0
                    cant=float(item_receta.get('Cantidad_Recursos', 1.0)); horas=float(item_receta.get('Cantidad_Horas', 1.0))
                    
                    monto_s = redondeo_excel(cant*horas*precio_s*factor, 3)
                    monto_d = redondeo_excel(monto_s/tc_inc_ce1 if tc_inc_ce1 > 0 else 0, 3)
                    
                    items_ce1.append({"descripcion": desc_item, "cantidad": cant, "horas": horas, "precio_soles": precio_s, "precio_dolares": round(precio_s / tc_inc_ce1 if tc_inc_ce1 > 0 else 0, 3), "factor_ajuste": factor, "monto_soles": monto_s, "monto_dolares": monto_d, "id_anexo": costo_final.get('ID_Anexo_Drive')})
                # (Finalizar fuentes y devolver - indentación correcta)
                if 'fuente_coti' in fuentes_ce1: fuentes_ce1['fuente_coti'] = "\n".join(filter(None, set(fuentes_ce1['fuente_coti'])))
                res_int['items_calculados'] = items_ce1
                res_int['fuentes'] = fuentes_ce1
                res_int['fuentes']['fi_mes'] = format_date(fecha_final_dt_ce1, "MMMM 'de' yyyy", locale='es')
                res_int['fuentes']['fi_ipc'] = float(ipc_inc_ce1)
                res_int['fuentes']['fi_tc'] = float(tc_inc_ce1)
            # (Bloque except de _calcular_ce1_interno - indentación correcta)
            except Exception as e_int:
                res_int['error'] = f"Error interno CE1: {e_int}"
            return res_int
        # --- FIN Lógica interna para calcular CE1 ---

        # (Llamada a _calcular_ce1_interno - indentación correcta)
        res_ce1 = _calcular_ce1_interno(datos_comunes, fecha_calculo_ce1)
        if res_ce1.get('error'):
            result['error'] = f"CE1: {res_ce1['error']}"
            return result
        # (Guardar resultados CE1 - indentación correcta)
        result['ce1_data_raw'] = res_ce1.get('items_calculados', [])
        result['ce1_soles'] = sum(item.get('monto_soles', 0) for item in result['ce1_data_raw'])
        result['ce1_dolares'] = sum(item.get('monto_dolares', 0) for item in result['ce1_data_raw'])
        result['ids_anexos'].update(item.get('id_anexo') for item in result['ce1_data_raw'] if item.get('id_anexo'))
        result['fuentes']['ce1'] = res_ce1.get('fuentes', {})

        # --- b. Calcular CE2 (Capacitación) ---
        res_ce2 = {}
        # *** FECHA DE CÁLCULO CE2: Siempre fecha_incumplimiento_extremo ***
        fecha_calculo_ce2 = fecha_final_dt # Usar el datetime object

        # (Condición if num_personal_ce2 - indentación correcta)
        if num_personal_ce2 > 0:
            datos_comunes_ce2 = {**datos_comunes, 'fecha_incumplimiento': fecha_calculo_ce2} # Usar fecha incumplimiento
            res_ce2 = calcular_costo_capacitacion(num_personal_ce2, datos_comunes_ce2)

            # (Validar error CE2 - indentación correcta)
            if res_ce2.get('error'):
                result['error'] = f"CE2: {res_ce2['error']}" # Error en CE2 siempre es fatal
                return result
            # (Guardar resultados CE2 - indentación correcta)
            elif res_ce2:
                result['ce2_data_raw'] = res_ce2.get('items_calculados', [])
                result['ce2_soles_calculado'] = sum(item.get('monto_soles', 0) for item in result['ce2_data_raw'])
                result['ce2_dolares_calculado'] = sum(item.get('monto_dolares', 0) for item in result['ce2_data_raw'])
                result['ids_anexos'].update(res_ce2.get('ids_anexos', []))
                result['fuentes']['ce2'] = {
                    'fuente_salario': res_ce2.get('fuente_salario', ''), 'pdf_salario': res_ce2.get('pdf_salario', ''),
                    'fuente_coti': res_ce2.get('fuente_coti', ''), 'fi_mes': res_ce2.get('fi_mes', ''),
                    'fi_ipc': res_ce2.get('fi_ipc', 0.0), 'fi_tc': res_ce2.get('fi_tc', 0.0)
                }

        # --- c. Aplicar Lógica Condicional para BI ---
        # (Asignación aplicar_ce2_a_bi - indentación correcta)
        result['aplicar_ce2_a_bi'] = (tipo_incumplimiento == "No remitió")

        # (Cálculo totales para BI - indentación correcta)
        result['ce_soles_para_bi'] = result['ce1_soles']
        result['ce_dolares_para_bi'] = result['ce1_dolares']
        if result['aplicar_ce2_a_bi']:
            result['ce_soles_para_bi'] += result['ce2_soles_calculado'] # Suma el CE2 calculado
            result['ce_dolares_para_bi'] += result['ce2_dolares_calculado']

        # (Limpiar error si todo ok - indentación correcta)
        if not result['error']: result['error'] = None

        return result

    # (Bloque except principal - indentación correcta)
    except Exception as e:
        import traceback; traceback.print_exc()
        result['error'] = f"Error crítico en _calcular_costo_evitado_extremo_inf005: {e}"
        return result

# ---------------------------------------------------------------------
# FUNCIÓN 2: RENDERIZAR INPUTS (CE2 Global)
# ---------------------------------------------------------------------

def renderizar_inputs_especificos(i, df_dias_no_laborables=None):
    st.markdown("##### Detalles de la Infracción (INF005)")
    datos_hecho = st.session_state.imputaciones_data[i]

    # --- INICIO CAMBIO 1: Autorelleno de Personal ---
    if 'tabla_personal' not in datos_hecho or not isinstance(datos_hecho['tabla_personal'], list):
        datos_hecho['tabla_personal'] = [
            {
                'Perfil': 'Gerente General', 
                'Descripción': (
                    "El Gerente General se encuentra encargado de:\n"
                    "- Planear, dirigir y aprobar los objetivos y metas inherentes a las actividades administrativas, operativas y financieras de la empresa.\n"
                    "- Administrar los recursos materiales, económicos y tecnológicos de la empresa.\n"
                    "- Supervisar, monitorear y evaluar el desarrollo de los procesos y sistemas que se lleva a cabo en la empresa.\n"
                    "Planificar, organizar y mantener canales de comunicación que garanticen la aplicación de las disposiciones necesarias para el cumplimiento de los objetivos de la empresa."
                ), 
                'Cantidad': 1
            },
            {
                'Perfil': 'Jefe de Seguridad, Salud Ocupacional y Medio Ambiente', 
                'Descripción': ("El Jefe de Seguridad, Salud Ocupacional y Medio Ambiente se encuentra encargado de:\n"
                                "- Implementar y gestionar el Sistema de Seguridad, Salud Ocupacional y Medio Ambiente y supervisar la correcta ejecución de las políticas, planes y actividades establecidas en el marco de la legislación vigente.\n"
                                "- Elaborar el Plan y Programa Anual de Seguridad, Salud y Medio ambiente en el Trabajo, según la normativa vigente.\n"
                                "- Elaborar el programa anual de entrenamiento y capacitación en temas de seguridad, salud y medio ambiente.\n"
                                "- Diseñar, implementar y liderar la ejecución del plan de auditorías e inspecciones tanto internas como externas en materia de Seguridad, Salud Ocupacional y Medio Ambiente."
                                ), 
                'Cantidad': 1
            }
        ]
    # --- FIN CAMBIO 1 ---

    st.markdown("###### **Extremos del incumplimiento**")
    if 'extremos' not in datos_hecho: datos_hecho['extremos'] = []

    if st.button("➕ Añadir Extremo", key=f"add_extremo_{i}"):
        datos_hecho['extremos'].append({})
        st.rerun()

    for j, extremo in enumerate(datos_hecho['extremos']):
        with st.container(border=True):
            st.markdown(f"**Extremo n.° {j + 1}**")

            # Inputs de Tipo y Periodo
            col_desc1, col_desc2 = st.columns(2)
            with col_desc1:
                extremo['tipo_monitoreo'] = st.text_input("Tipo de Informe/Monitoreo", key=f"tipo_monitoreo_{i}_{j}", value=extremo.get('tipo_monitoreo', ''), placeholder="Ej: Monitoreo de Calidad de Aire")
            with col_desc2:
                extremo['periodicidad'] = st.text_input("Periodicidad/Periodo", key=f"periodicidad_{i}_{j}", value=extremo.get('periodicidad', ''), placeholder="Ej: Cuarto Trimestre 2021")

            # Inputs de Fechas e Incumplimiento
            col_fecha_max, col_fecha_inc, col_tipo = st.columns(3)
            with col_fecha_max:
                fecha_max_input = st.date_input("Fecha máxima de presentación", key=f"fecha_max_{i}_{j}", value=extremo.get('fecha_maxima_presentacion'), format="DD/MM/YYYY", max_value=date.today())
                extremo['fecha_maxima_presentacion'] = fecha_max_input
            with col_fecha_inc:
                if fecha_max_input:
                    fecha_inc_calculada = fecha_max_input + timedelta(days=1)
                    extremo['fecha_incumplimiento'] = fecha_inc_calculada
                    st.metric("Fecha de incumplimiento", fecha_inc_calculada.strftime('%d/%m/%Y'))
                else:
                    extremo['fecha_incumplimiento'] = None
            with col_tipo:
                tipo_extremo = st.radio("Tipo de incumplimiento", ["No remitió", "Remitió fuera de plazo"], key=f"tipo_extremo_{i}_{j}", index=0 if extremo.get('tipo_extremo') == "No remitió" else 1 if extremo.get('tipo_extremo') == "Remitió fuera de plazo" else None, horizontal=True)
                extremo['tipo_extremo'] = tipo_extremo

            if tipo_extremo == "Remitió fuera de plazo":
                fecha_inc_actual = extremo.get('fecha_incumplimiento')
                min_fecha_ext = fecha_inc_actual if fecha_inc_actual else date.today()
                extremo['fecha_extemporanea'] = st.date_input("Fecha de cumplimiento extemporáneo", min_value=min_fecha_ext, key=f"fecha_ext_{i}_{j}", value=extremo.get('fecha_extemporanea'), format="DD/MM/YYYY")
            else:
                extremo['fecha_extemporanea'] = None

            # --- INICIO CAMBIO 2: Tabla de Personal integrada dentro del Extremo ---
            st.divider()
            st.markdown("###### **Personal a capacitar (CE2)**")
            df_personal = pd.DataFrame(datos_hecho['tabla_personal'])

            if j == 0:
                # Solo editable en el primer bloque para evitar duplicidad de llaves
                edited_df = st.data_editor(
                    df_personal,
                    num_rows="dynamic",
                    key=f"data_editor_personal_{i}_{j}", 
                    hide_index=True,
                    use_container_width=True,
                    column_config={
                        "Perfil": st.column_config.TextColumn("Perfil", required=True),
                        "Descripción": st.column_config.TextColumn("Descripción", width="large"),
                        "Cantidad": st.column_config.NumberColumn("Cantidad", min_value=0, step=1, required=True, format="%d"),
                    }
                )
                datos_hecho['tabla_personal'] = edited_df.to_dict('records')
            else:
                # Solo lectura en los siguientes bloques
                st.dataframe(df_personal, use_container_width=True, hide_index=True)

            # Cálculo de totales para el sistema
            cant_num = [pd.to_numeric(p.get('Cantidad'), errors='coerce') for p in datos_hecho['tabla_personal']]
            total_pers = int(pd.Series(cant_num).fillna(0).sum())
            datos_hecho['num_personal_capacitacion'] = total_pers
            
            if j == 0:
                st.metric("Total de Personal a Capacitar", f"{total_pers}")
            # --- FIN CAMBIO 2 ---

            if st.button(f"🗑️ Eliminar Extremo {j + 1}", key=f"del_extremo_{i}_{j}"):
                datos_hecho['extremos'].pop(j)
                st.rerun()

    return datos_hecho

# ---------------------------------------------------------------------
# FUNCIÓN 3: VALIDACIÓN DE INPUTS (CE2 Global)
# ---------------------------------------------------------------------

def validar_inputs(datos_hecho):
    """
    Valida inputs de INF005 con CE2 global.
    Verifica que el total de personal sea mayor a 0 y que
    todos los campos requeridos en cada extremo estén completos.
    """
    # 1. Validar CE2 (Global)
    # Checks if the total number of personnel to be trained is greater than 0.
    if not datos_hecho.get('num_personal_capacitacion', 0) > 0:
        # st.warning("Debe ingresar al menos una persona en la tabla de 'Personal a capacitar'.") # Optional warning message
        return False # Fails validation if no personnel specified

    # 2. Validar Extremos
    # Checks if the list of 'extremos' exists and is not empty.
    if not datos_hecho.get('extremos'):
        # st.warning("Debe añadir al menos un extremo de incumplimiento.") # Optional warning message
        return False # Fails validation if no extremes are added

    # Loop through each 'extremo' dictionary in the 'extremos' list.
    for j, extremo in enumerate(datos_hecho.get('extremos', [])):
        # Validar campos básicos del extremo
        # Check if all essential fields for the extreme have been filled.
        if not all([
            extremo.get('tipo_monitoreo'),         # Check for report type
            extremo.get('periodicidad'),          # Check for period/frequency
            extremo.get('fecha_maxima_presentacion'), # Check for the deadline date input by the user
            extremo.get('tipo_extremo')             # Check if non-compliance type is selected
        ]):
            # st.warning(f"Complete campos básicos del Extremo n.° {j + 1}.") # Optional warning message
            return False # Fails validation if any basic field is missing

        # Validar campo condicional (fecha extemporánea)
        # If the non-compliance type is 'Remitió fuera de plazo', check if the late submission date is provided.
        if extremo.get('tipo_extremo') == "Remitió fuera de plazo" and not extremo.get('fecha_extemporanea'):
            # st.warning(f"Ingrese 'Fecha de cumplimiento extemporáneo' en Extremo n.° {j + 1}.") # Optional warning message
            return False # Fails validation if late date is missing when required

    # Si todas las validaciones pasan
    return True # Returns True if all checks pass

# ---------------------------------------------------------------------
# FUNCIÓN 4: DESPACHADOR PRINCIPAL (Sin cambios respecto a la versión anterior)
# ---------------------------------------------------------------------
def procesar_infraccion(datos_comunes, datos_hecho):
    """
    Decide si procesar como hecho simple (1 extremo) o múltiple (>1 extremo)
    y llama a la función correspondiente.
    """
    # Cuenta cuántos diccionarios (extremos) hay en la lista 'extremos'
    # Si 'extremos' no existe o está vacío, devuelve 0.
    num_extremos = len(datos_hecho.get('extremos', []))

    # Validación básica: Si no se ingresó ningún extremo, devuelve un error.
    if num_extremos == 0:
        return {'error': 'No se ha registrado ningún extremo para este hecho.'}
    # Si hay exactamente un extremo, llama a la función _procesar_hecho_simple.
    elif num_extremos == 1:
        return _procesar_hecho_simple(datos_comunes, datos_hecho)
    # Si hay más de un extremo, llama a la función _procesar_hecho_multiple.
    else: # num_extremos > 1
        return _procesar_hecho_multiple(datos_comunes, datos_hecho)

# --- Asegúrate que las funciones _procesar_hecho_simple y _procesar_hecho_multiple
# --- estén definidas DESPUÉS de esta función en tu archivo INF005.py ---

# ---------------------------------------------------------------------
# FUNCIÓN 5: PROCESAR HECHO SIMPLE (Usa CE unificado)
# ---------------------------------------------------------------------
def _procesar_hecho_simple(datos_comunes, datos_hecho):
    """
    Procesa un hecho INF005 con un único extremo, usando UNA SOLA plantilla
    que contiene un bloque condicional {% if aplicar_capacitacion %}.
    """
    try:
        # 1. Obtener datos del extremo y calcular CE (Sin cambios)
        extremo = datos_hecho['extremos'][0]
        fecha_inc = extremo.get('fecha_incumplimiento')
        tipo_inc = extremo.get('tipo_extremo')
        fecha_ext = extremo.get('fecha_extemporanea')

        res_ce = _calcular_costo_evitado_extremo_inf005(datos_comunes, datos_hecho, extremo)
        if res_ce.get('error'):
            return {'error': f"Error al calcular CE: {res_ce['error']}"}

        # Determinar si aplica capacitación (Sin cambios)
        # Esta variable controlará el {% if %} en la plantilla
        aplicar_capacitacion = res_ce['aplicar_ce2_a_bi']

        # --- INICIO CAMBIO PLANTILLAS ---
        # 2. Cargar UNA SOLA plantilla BI y la de Anexo CE
        df_tipificacion, id_infraccion = datos_comunes['df_tipificacion'], datos_comunes['id_infraccion']
        filas_inf = df_tipificacion[df_tipificacion['ID_Infraccion'] == id_infraccion]
        if filas_inf.empty:
            return {'error': f"No se encontró ID '{id_infraccion}' en Tipificación."}
        fila_inf = filas_inf.iloc[0]

        # Obtener ID de la plantilla UNIFICADA (la que tiene el {% if %})
        # **DEBES ASEGURARTE** que la columna 'ID_Plantilla_BI_ConCap' (o la que elijas)
        # en tu Google Sheet contenga el ID de la plantilla que tiene el {% if %}
        id_tpl_unificada = fila_inf.get('ID_Plantilla_BI') # O cambia esto al nombre de columna correcto
        id_tpl_ce_anexo = fila_inf.get('ID_Plantilla_CE') # Anexo CE se mantiene único

        # Validar IDs necesarios
        if not id_tpl_unificada or not id_tpl_ce_anexo:
             # Ajusta el mensaje de error para reflejar que buscas una plantilla unificada
             return {'error': f'Faltan IDs de plantilla BI Unificada (ej: ID_Plantilla_BI_ConCap) o Anexo CE para {id_infraccion}.'}

        # Descargar las plantillas
        buf_bi = descargar_archivo_drive(id_tpl_unificada)
        buf_ce = descargar_archivo_drive(id_tpl_ce_anexo)
        if not buf_bi or not buf_ce:
            return {'error': f'Fallo descarga plantilla BI {id_tpl_unificada} o anexo {id_tpl_ce_anexo}.'}

        # Crear objetos DocxTemplate
        doc_tpl_bi = DocxTemplate(buf_bi) # Plantilla BI única con {% if %}
        tpl_anexo = DocxTemplate(buf_ce) # Plantilla Anexo CE
        # --- FIN CAMBIO PLANTILLAS ---

        # 3. Calcular Beneficio Ilícito (BI) y Multa
        
        # --- CORRECCIÓN 2: Usar el texto del Hecho Imputado (UI) ---
        # Toma lo que escribiste en el cuadro de texto de la app
        texto_bi_preciso = datos_hecho.get('texto_hecho', '')
        # -----------------------------------------------------

        datos_bi_base = {
            **datos_comunes,
            'ce_soles': res_ce['ce_soles_para_bi'],
            'ce_dolares': res_ce['ce_dolares_para_bi'],
            'fecha_incumplimiento': fecha_inc,
            'texto_del_hecho': texto_bi_preciso # <-- USAR LA VARIABLE PRECISA
        }
        # --- DEFINIR VARIABLE AQUÍ ---
        es_extemporaneo = (tipo_inc == "Remitió fuera de plazo")
        
        res_bi = None
        if es_extemporaneo:
            pre_bi = calcular_beneficio_ilicito(datos_bi_base);
            if pre_bi.get('error'): return pre_bi
            datos_bi_ext = {**datos_bi_base, 'fecha_cumplimiento_extemporaneo': fecha_ext, **pre_bi}
            res_bi = calcular_beneficio_ilicito_extemporaneo(datos_bi_ext)
        else: res_bi = calcular_beneficio_ilicito(datos_bi_base)
        if not res_bi or res_bi.get('error'):
            return res_bi or {'error': 'Error desconocido al calcular el BI.'}
        bi_uit = res_bi.get('beneficio_ilicito_uit', 0)

        # --- INICIO: Lógica de Moneda (Basado en INF004) ---
        moneda_calculo = res_bi.get('moneda_cos', 'USD') 
        es_dolares = (moneda_calculo == 'USD')
        
        if es_dolares:
            texto_moneda_bi = "moneda extranjera (Dólares)"
            ph_bi_abreviatura_moneda = "US$"
        else:
            texto_moneda_bi = "moneda nacional (Soles)"
            ph_bi_abreviatura_moneda = "S/"
        # --- FIN: Lógica de Moneda ---

        # --- INICIO: Recuperar Factor F y Calcular Multa ---
        # --- INICIO: Recuperar Factor F y Calcular Multa ---
        factor_f = datos_hecho.get('factor_f_calculado', 1.0)

        res_multa = calcular_multa({
            **datos_comunes, 
            'beneficio_ilicito': bi_uit,
            'factor_f': factor_f # <--- NUEVO
        })
        multa_uit = res_multa.get('multa_final_uit', 0)
        # --- FIN ---

        # --- INICIO: LÓGICA DE REDUCCIÓN Y TOPE (Simple) ---
        datos_hecho_completos = datos_comunes.get('datos_hecho_completos', {})
        aplica_reduccion_str = datos_hecho_completos.get('aplica_reduccion', 'No')
        porcentaje_str = datos_hecho_completos.get('porcentaje_reduccion', '0%')
        multa_con_reduccion_uit = multa_uit
        
        if aplica_reduccion_str == 'Sí':
            reduccion_factor = 0.5 if porcentaje_str == '50%' else 0.7
            multa_con_reduccion_uit = redondeo_excel(multa_uit * reduccion_factor, 3)

        infraccion_info = datos_comunes['df_tipificacion'][datos_comunes['df_tipificacion']['ID_Infraccion'] == id_infraccion]
        tope_multa_uit = float('inf')
        if not infraccion_info.empty and pd.notna(infraccion_info.iloc[0].get('Tope_Multa_Infraccion')):
            tope_multa_uit = float(infraccion_info.iloc[0]['Tope_Multa_Infraccion'])

        multa_final_del_hecho_uit = min(multa_con_reduccion_uit, tope_multa_uit)
        se_aplica_tope = multa_con_reduccion_uit > tope_multa_uit
        multa_reducida_uit = multa_con_reduccion_uit if aplica_reduccion_str == 'Sí' else multa_uit
        # --- FIN: LÓGICA DE REDUCCIÓN Y TOPE ---

        # --- Definir variables condicionales para la plantilla ---
        if aplicar_capacitacion:
            label_ce_principal = "CE1"  # Define the label when capacitación applies
            show_ce2_block = True      # Flag to show the CE2 block
            # (You would still need to handle the footnote references [1] & [2] somehow)
        else:
            label_ce_principal = "CE"   # Define the label when capacitación DOES NOT apply
            show_ce2_block = False     # Flag to hide the CE2 block
            # (You would still need to handle the footnote references [3] & [4] somehow)

        # 4. Generar Tablas Subdocumento para el CUERPO del informe (Sin cambios en cómo se generan)
        # Se generan igual, pero solo se usarán en la plantilla si el 'if' es True
        ce1_fmt = []
        for item in res_ce['ce1_data_raw']:
            desc_orig = item.get('descripcion', '')
            texto_adicional = ""
            if "Profesional" in desc_orig: texto_adicional = "1/ "
            elif "Alquiler de laptop" in desc_orig: texto_adicional = "2/ "
            
            ce1_fmt.append({
                'descripcion': f"{desc_orig}{texto_adicional}",
                'cantidad': format_decimal_dinamico(item.get('cantidad', 0)),
                'horas': format_decimal_dinamico(item.get('horas', 0)),
                'precio_soles': f"S/ {item.get('precio_soles', 0):,.3f}",
                'factor_ajuste': f"{item.get('factor_ajuste', 0):,.3f}",
                'monto_soles': f"S/ {item.get('monto_soles', 0):,.3f}",
                'monto_dolares': f"US$ {item.get('monto_dolares', 0):,.3f}"
            })
        
        ce1_fmt.append({
            'descripcion': 'Total',
            'monto_soles': f"S/ {res_ce['ce1_soles']:,.3f}",
            'monto_dolares': f"US$ {res_ce['ce1_dolares']:,.3f}"
        })

        tabla_ce1 = create_table_subdoc(
            doc_tpl_bi, 
            ["Descripción", "Cantidad", "Horas", "Precio asociado (S/)", "Factor de ajuste 3/", "Monto (*) (S/)", "Monto (*) (US$) 4/"],
            ce1_fmt, 
            ['descripcion', 'cantidad', 'horas', 'precio_soles', 'factor_ajuste', 'monto_soles', 'monto_dolares']
        )

        tabla_ce2 = None # Inicializa como None
        ce2_fmt = []
        if res_ce['ce2_data_raw']:
            for item in res_ce['ce2_data_raw']:
                ce2_fmt.append({
                    'descripcion': f"{item.get('descripcion', '')} 1/",
                    'precio_dolares': f"US$ {item.get('precio_dolares', 0):,.3f}",
                    'precio_soles': f"S/ {item.get('precio_soles', 0):,.3f}",
                    'factor_ajuste': f"{item.get('factor_ajuste', 0):,.3f}",
                    'monto_soles': f"S/ {item.get('monto_soles', 0):,.3f}",
                    'monto_dolares': f"US$ {item.get('monto_dolares', 0):,.3f}"
                })
            
            ce2_fmt.append({
                'descripcion': 'Total',
                'monto_soles': f"S/ {res_ce['ce2_soles_calculado']:,.3f}",
                'monto_dolares': f"US$ {res_ce['ce2_dolares_calculado']:,.3f}"
            })

            tabla_ce2 = create_table_subdoc(
                doc_tpl_bi, 
                ["Descripción", "Precio (US$)", "Precio (S/)", "Factor de ajuste 2/", "Monto (*) (S/)", "Monto (*) (US$) 3/"],
                ce2_fmt, 
                ['descripcion', 'precio_dolares', 'precio_soles', 'factor_ajuste', 'monto_soles', 'monto_dolares']
            )

        # --- SOLUCIÓN: Compactar y Reordenar Notas al Pie de BI (Basado en INF004) ---
        filas_bi_crudas = res_bi.get('table_rows', [])
        fn_map_orig = res_bi.get('footnote_mapping', {})
        fn_data = res_bi.get('footnote_data', {})
        
        # 1. Identificar letras realmente usadas en esta tabla
        letras_usadas = sorted(list({r for f in filas_bi_crudas if f.get('ref') for r in f.get('ref').replace(" ", "").split(",") if r}))
        
        # 2. Crear mapeo secuencial (a, b, c...)
        letras_base = "abcdefghijklmnopqrstuvwxyz"
        map_traduccion = {v: letras_base[i] for i, v in enumerate(letras_usadas)}
        nuevo_fn_map = {map_traduccion[v]: fn_map_orig[v] for v in letras_usadas if v in fn_map_orig}

        # 3. Re-etiquetar filas de la tabla combinando superíndices
        filas_bi_para_tabla = []
        for fila in filas_bi_crudas:
            nueva_fila = fila.copy()
            ref_orig = nueva_fila.get('ref', '')
            super_final = str(nueva_fila.get('descripcion_superindice', ''))
            if ref_orig:
                nuevas = [map_traduccion[r] for r in ref_orig.replace(" ", "").split(",") if r in map_traduccion]
                if nuevas: super_final += f"({', '.join(nuevas)})"
            
            filas_bi_para_tabla.append({
                'descripcion_texto': fila.get('descripcion_texto', ''),
                'descripcion_superindice': super_final,
                'monto': fila.get('monto', '')
            })

        # 4. Generar lista de notas filtrada y en orden
        # 4. Generar lista de notas filtrada y en orden (CORREGIDO es_extemporaneo)
        fn_list = [f"({l}) {obtener_fuente_formateada(k, fn_data, id_infraccion, es_extemporaneo)}" for l, k in sorted(nuevo_fn_map.items())]
        fn_data_dict = {'list': fn_list, 'elaboration': 'Elaboración: Subdirección de Sanción y Gestión de Incentivos (SSAG) - DFAI.', 'style': 'FuenteTabla'}
        
        tabla_bi = create_main_table_subdoc(doc_tpl_bi, ["Descripción", "Monto"], filas_bi_para_tabla, keys=['descripcion_texto', 'monto'], footnotes_data=fn_data_dict, column_widths=(5, 1))
        tabla_multa = create_main_table_subdoc(doc_tpl_bi, ["Componentes", "Monto"], res_multa.get('multa_data_raw', []), ['Componentes', 'Monto'], texto_posterior="Elaboración: Subdirección de Sanción y Gestión de Incentivos (SSAG) - DFAI.", estilo_texto_posterior='FuenteTabla', column_widths=(5, 1))

        tabla_personal = None # Inicializa como None
        tabla_pers_data = [] # Para la app
        if aplicar_capacitacion: # Solo crea la tabla si aplica capacitación
            tabla_pers_render = datos_hecho.get('tabla_personal', [])
            tabla_pers_sin_total = []
            for fila in tabla_pers_render:
                perfil = fila.get('Perfil'); cantidad = pd.to_numeric(fila.get('Cantidad'), errors='coerce')
                if perfil and cantidad > 0: tabla_pers_sin_total.append({'Perfil': perfil, 'Descripción': fila.get('Descripción', ''), 'Cantidad': int(cantidad)})
            num_pers_total_int = int(datos_hecho.get('num_personal_capacitacion', 0))
            tabla_pers_data = tabla_pers_sin_total + [{'Perfil':'Total', 'Descripción':'', 'Cantidad': num_pers_total_int}]
            tabla_personal = create_personal_table_subdoc(doc_tpl_bi, ["Perfil", "Descripción", "Cantidad"], tabla_pers_data, ['Perfil', 'Descripción', 'Cantidad'], column_widths=(2,3,1), texto_posterior="Elaboración: Subdirección de Sanción y Gestión de Incentivos (SSAG) - DFAI.")


        # 5. Contexto y Renderizado del CUERPO
        fuentes_ce = res_ce.get('fuentes', {})
        contexto_word = {
            **datos_comunes['context_data'],
            'acronyms': datos_comunes['acronym_manager'],
            'hecho': {
                'numero_imputado': datos_comunes['numero_hecho_actual'],
                'descripcion': RichText(datos_hecho.get('texto_hecho', '')) # Usa RichText para la descripción
            },
            'numeral_hecho': f"IV.{datos_comunes['numero_hecho_actual'] + 1}", # Suma 1 para empezar en IV.2

            # --- INICIO ADICIÓN: Placeholders de Tipo y Periodo ---
            'tipo_monitoreo': extremo.get('tipo_monitoreo', ''),
            'periodicidad': extremo.get('periodicidad', ''),
            # --- FIN ADICIÓN ---

            # --- INICIO ADICIÓN: Fechas Largas ---
            'fecha_incumplimiento_larga': (format_date(fecha_inc, "d 'de' MMMM 'de' yyyy", locale='es') if fecha_inc else "N/A").replace("septiembre", "setiembre").replace("Septiembre", "Setiembre"),
            'fecha_extemporanea_larga': (format_date(fecha_ext, "d 'de' MMMM 'de' yyyy", locale='es') if fecha_ext else "N/A").replace("septiembre", "setiembre").replace("Septiembre", "Setiembre"),
            # --- FIN ADICIÓN ---

            # --- PASAR LA CONDICIÓN A LA PLANTILLA ---
            'aplicar_capacitacion': aplicar_capacitacion, # <-- La variable booleana para el {% if %}
            'label_ce_principal': label_ce_principal,
            # --- FIN ---

            'tabla_ce1': tabla_ce1,
            'tabla_ce2': tabla_ce2, # Pasas la tabla CE2 (puede ser None)
            'tabla_bi': tabla_bi,
            'tabla_multa': tabla_multa,
            'tabla_detalle_personal': tabla_personal, # Pasas la tabla de personal (puede ser None)
            'num_personal_total_texto': texto_con_numero(datos_hecho.get('num_personal_capacitacion', 0), 'f'), # Pasas el texto del número total
            'mh_uit': f"{multa_uit:,.3f} UIT",
            'bi_uit': f"{bi_uit:,.3f} UIT",
            'bi_moneda_es_dolares': es_dolares,
            'ph_bi_moneda_texto': texto_moneda_bi,
            'ph_bi_moneda_simbolo': ph_bi_abreviatura_moneda,
            'bi_moneda_es_soles': (moneda_calculo == 'PEN'),
            # --- INICIO: PLACEHOLDERS DE REDUCCIÓN Y TOPE ---
            'aplica_reduccion': aplica_reduccion_str == 'Sí',
            'porcentaje_reduccion': porcentaje_str,
            'texto_reduccion': datos_hecho_completos.get('texto_reduccion', ''),
            'memo_num': datos_hecho_completos.get('memo_num', ''),
            'memo_fecha': format_date(datos_hecho_completos.get('memo_fecha'), "d 'de' MMMM 'de' yyyy", locale='es') if datos_hecho_completos.get('memo_fecha') else '',
            'escrito_num': datos_hecho_completos.get('escrito_num', ''),
            'escrito_fecha': format_date(datos_hecho_completos.get('escrito_fecha'), "d 'de' MMMM 'de' yyyy", locale='es') if datos_hecho_completos.get('escrito_fecha') else '',
            'multa_con_reduccion_uit': f"{multa_con_reduccion_uit:,.3f} UIT",
            'se_aplica_tope': se_aplica_tope,
            'tope_multa_uit': f"{tope_multa_uit:,.3f} UIT",
            # --- FIN: PLACEHOLDERS DE REDUCCIÓN Y TOPE ---
             'fuente_cos': res_bi.get('fuente_cos', ''), # Añadido por si falta
            **(fuentes_ce.get('ce1', {}).get('placeholders_dinamicos', {})),
            # --- Texto Explicativo (Se pasa vacío, se genera en app.py) ---
            'texto_explicacion_prorrateo': '',
            'ph_ipc_promedio_salario_ce1': fuentes_ce.get('ce1', {}).get('texto_ipc_costeo_salario', ''),
        }
        # Renderizar la plantilla única con el contexto
        doc_tpl_bi.render(contexto_word, autoescape=True)
        buf_final_hecho = io.BytesIO()
        doc_tpl_bi.save(buf_final_hecho)

        # 6. Generar Anexo CE (Simple) (Sin cambios aparentes necesarios aquí)
        anexos_ce = []
        # Reutilizamos ce1_fmt que ya tiene el formato y el total
        tabla_ce1_anx = create_table_subdoc(
            tpl_anexo, 
            ["Descripción", "Cantidad", "Horas", "Precio asociado (S/)", "Factor de ajuste 3/", "Monto (*) (S/)", "Monto (*) (US$) 4/"],
            ce1_fmt, 
            ['descripcion', 'cantidad', 'horas', 'precio_soles', 'factor_ajuste', 'monto_soles', 'monto_dolares']
        )
        tabla_ce2_anx = None
        if res_ce['ce2_data_raw']: # Solo crea tabla anexo CE2 si hay datos
             # Reutilizamos ce2_fmt que ya tiene el formato y el total
             tabla_ce2_anx = create_table_subdoc(
                tpl_anexo, 
                ["Descripción", "Precio (US$)", "Precio (S/)", "Factor de ajuste 2/", "Monto (*) (S/)", "Monto (*) (US$) 3/"],
                ce2_fmt, 
                ['descripcion', 'precio_dolares', 'precio_soles', 'factor_ajuste', 'monto_soles', 'monto_dolares']
            )
        # --- Crear tabla resumen ANEXO (igual que antes) ---
        resumen_anexo = [{'desc': 'Costo de sistematización y remisión de la información - CE1', 'sol': f"S/ {res_ce['ce1_soles']:,.3f}", 'dol': f"US$ {res_ce['ce1_dolares']:,.3f}"}]
        if aplicar_capacitacion: # Condición para añadir CE2 al resumen del anexo
            resumen_anexo.append({'desc': 'Costo de capacitación al personal - CE2', 'sol': f"S/ {res_ce['ce2_soles_calculado']:,.3f}", 'dol': f"US$ {res_ce['ce2_dolares_calculado']:,.3f}"})
        resumen_anexo.append({'desc': 'Costo Evitado Total', 'sol': f"S/ {res_ce['ce_soles_para_bi']:,.3f}", 'dol': f"US$ {res_ce['ce_dolares_para_bi']:,.3f}"})
        tabla_resumen_anx = create_table_subdoc(tpl_anexo, ["Componente", "Monto (*) (S/)", "Monto (*) (US$)"], resumen_anexo, ['desc', 'sol', 'dol'])
        # ---
        contexto_anx = {
            **contexto_word, # Usar el mismo contexto base que el cuerpo
            'extremo': { # Datos específicos del extremo para el anexo
                 'tipo': extremo.get('tipo_monitoreo', ''),
                 'periodicidad': extremo.get('periodicidad', ''),
                 'tipo_incumplimiento': tipo_inc,
                 'fecha_incumplimiento': format_date(fecha_inc, "d/MM/yyyy"),
                 'fecha_extemporanea': format_date(fecha_ext, "d/MM/yyyy") if fecha_ext else "N/A",
            },
            'tabla_ce1_anexo': tabla_ce1_anx,
            'tabla_ce2_anexo': tabla_ce2_anx, # Pasas la tabla (puede ser None)
            'tabla_resumen_anexo': tabla_resumen_anx,
            # Pasar todas las fuentes necesarias para el anexo
            # --- INICIO CORRECCIÓN 1: Placeholders Unificados ---
            # Fuentes de CE1 (Remisión)
            'fuente_salario_ce1': fuentes_ce.get('ce1', {}).get('fuente_salario', ''),
            'pdf_salario_ce1': fuentes_ce.get('ce1', {}).get('pdf_salario', ''),
            'sustento_prof_ce1': fuentes_ce.get('ce1', {}).get('sustento_item_profesional', ''),
            'fuente_coti_ce1': fuentes_ce.get('ce1', {}).get('fuente_coti', ''),
            
            # Fuentes de CE2 (Capacitación)
            'fuente_salario_ce2': fuentes_ce.get('ce2', {}).get('fuente_salario', ''),
            'pdf_salario_ce2': fuentes_ce.get('ce2', {}).get('pdf_salario', ''),
            'fuente_coti_ce2': fuentes_ce.get('ce2', {}).get('fuente_coti', ''),
            # Fuentes Comunes (IPC/TC de la fecha de incumplimiento)
            'fi_mes': fuentes_ce.get('fi_mes', ''),
            'fi_ipc': f"{fuentes_ce.get('fi_ipc', 0)}",
            'fi_tc': f"{fuentes_ce.get('fi_tc', 0)}",
            'ph_ipc_promedio_salario_ce1': fuentes_ce.get('ce1', {}).get('texto_ipc_costeo_salario', ''),
            # --- FIN CORRECCIÓN 1 ---
        }
        tpl_anexo.render(contexto_anx, autoescape=True)
        buf_anexo_final = io.BytesIO()
        tpl_anexo.save(buf_anexo_final)
        anexos_ce.append(buf_anexo_final)

        # 7. Devolver Resultados (Ajustar flags y datos pasados a app.py)
        resultados_app = {
             'extremos': [{ # Mantener estructura para consistencia
                  'tipo': f"{extremo.get('tipo_monitoreo')} ({tipo_inc})",
                  'ce1_data': res_ce['ce1_data_raw'], 'ce2_data': res_ce['ce2_data_raw'],
                  'ce1_soles': res_ce['ce1_soles'], 'ce1_dolares': res_ce['ce1_dolares'], # <--- AÑADIDO
                  'ce2_soles_calculado': res_ce['ce2_soles_calculado'], 'ce2_dolares_calculado': res_ce['ce2_dolares_calculado'], # <--- AÑADIDO
                  'ce_soles_para_bi': res_ce['ce_soles_para_bi'], 'ce_dolares_para_bi': res_ce['ce_dolares_para_bi'], # <--- AÑADIDO
                  'bi_data': res_bi.get('table_rows', []), 'bi_uit': bi_uit,
                  'aplicar_ce2_a_bi': aplicar_capacitacion # Flag para UI
             }],
             'totales': { # Totales del hecho
                  'ce1_total_soles': res_ce['ce1_soles'], 'ce1_total_dolares': res_ce['ce1_dolares'], # <--- AÑADIDO
                  'ce2_total_soles_calculado': res_ce['ce2_soles_calculado'], 'ce2_total_dolares_calculado': res_ce['ce2_dolares_calculado'], # <--- AÑADIDO
                  'ce_total_soles_para_bi': res_ce['ce_soles_para_bi'], 'ce_total_dolares_para_bi': res_ce['ce_dolares_para_bi'], # <--- AÑADIDO
                  'beneficio_ilicito_uit': bi_uit, 'multa_final_uit': multa_uit,
                  'bi_data_raw': res_bi.get('table_rows', []),
                  'multa_data_raw': res_multa.get('multa_data_raw', []),
                  'tabla_personal_data': tabla_pers_data, # Datos tabla personal
                  'aplicar_ce2_a_bi': aplicar_capacitacion, # Flag para UI
                # --- CORRECCIÓN 1: Pasar variables de reducción correctas ---
                  'aplica_reduccion': aplica_reduccion_str,
                  'porcentaje_reduccion': porcentaje_str,
                  'multa_con_reduccion_uit': multa_con_reduccion_uit,
                  'multa_reducida_uit': multa_reducida_uit,
                  'multa_final_aplicada': multa_final_del_hecho_uit # <-- ESTA ES LA QUE USA APP.PY PARA LA TABLA
                  # ----------------------------------------------------------
             }      
        }
        return {
            # 'contexto_final_word': contexto_word, # Ya no necesitas pasar el contexto crudo
            'doc_pre_compuesto': buf_final_hecho, # Pasas el buffer ya renderizado
            'resultados_para_app': resultados_app,
            'es_extemporaneo': es_extemporaneo,
            'usa_capacitacion': aplicar_capacitacion, # Flag para app.py
            'anexos_ce_generados': anexos_ce,
            'ids_anexos': list(filter(None, res_ce.get('ids_anexos', set()))),
            'texto_explicacion_prorrateo': '', # No aplica en simple
            'tabla_detalle_personal': tabla_personal, # Subdoc tabla personal (puede ser None)
            'tabla_personal_data': tabla_pers_data # Datos tabla personal (puede ser [])
        }

    except Exception as e:
        import traceback; traceback.print_exc()
        try:
            import streamlit as st
            st.error(f"Error procesando hecho simple INF005 (plantilla única): {e}")
        except ImportError:
            print(f"Error procesando hecho simple INF005 (plantilla única): {e}")
        return {'error': f"Error procesando hecho simple INF005 (plantilla única): {e}"}
    

# ---------------------------------------------------------------------
# FUNCIÓN 6: PROCESAR HECHO MÚLTIPLE (Adaptado a INF005)
# ---------------------------------------------------------------------
def _procesar_hecho_multiple(datos_comunes, datos_hecho):
    """
    Procesa INF005 con múltiples extremos usando UNA SOLA plantilla principal
    con bucle {% for %} y condicionales {% if %}. Incluye las fuentes debajo de CADA tabla BI.
    """
    try:
        # --- 1. Cargar Plantillas (Sin cambios) ---
        df_tipificacion, id_infraccion = datos_comunes['df_tipificacion'], datos_comunes['id_infraccion']
        filas_inf = df_tipificacion[df_tipificacion['ID_Infraccion'] == id_infraccion]
        if filas_inf.empty: return {'error': f"No se encontró ID '{id_infraccion}' en Tipificación."}
        fila_inf = filas_inf.iloc[0]
        # Asume que ID_Plantilla_BI_Extremo es la plantilla con el bucle y los if
        id_tpl_principal = fila_inf.get('ID_Plantilla_BI_Extremo')
        id_tpl_anx = fila_inf.get('ID_Plantilla_CE_Extremo') # Anexo CE

        if not id_tpl_principal or not id_tpl_anx:
             return {'error': f'Faltan IDs de plantilla (BI_Extremo o CE_Extremo) para {id_infraccion}.'}
        buffer_plantilla = descargar_archivo_drive(id_tpl_principal)
        buffer_anexo = descargar_archivo_drive(id_tpl_anx)
        if not buffer_plantilla or not buffer_anexo:
            return {'error': f'Fallo descarga plantilla BI Extremo {id_tpl_principal} o anexo {id_tpl_anx}.'}
        tpl_principal = DocxTemplate(buffer_plantilla)
        # --- FIN CARGA PLANTILLAS ---

        # 2. Inicializar acumuladores (Sin cambios)
        total_bi_uit = 0.0
        lista_bi_resultados_completos = [] # Guardar resultados BI para multa final
        anexos_ids = set()
        num_hecho = datos_comunes['numero_hecho_actual']
        anexos_ce = []
        lista_extremos_plantilla_word = []
        # Determinar si la capacitación aplica en general (para tabla personal y texto explicativo)
        aplicar_capacitacion_general = any(ext.get('tipo_extremo') == 'No remitió' for ext in datos_hecho['extremos'])
        resultados_app = { 'extremos': [], 'totales': {'ce1_total_soles': 0.0, 'ce2_total_soles_calculado': 0.0, 'ce_total_soles_para_bi': 0.0, 'aplicar_ce2_a_bi': aplicar_capacitacion_general} }

        # 3. Generar Tabla Personal (si aplica, una sola vez antes del bucle) (Sin cambios)
        tabla_pers_subdoc_final = None
        tabla_pers_data = []
        num_pers_total_int = 0
        if aplicar_capacitacion_general:
            tabla_pers_render = datos_hecho.get('tabla_personal', [])
            tabla_pers_sin_total = []
            for fila in tabla_pers_render:
                perfil = fila.get('Perfil'); cantidad = pd.to_numeric(fila.get('Cantidad'), errors='coerce')
                if perfil and cantidad > 0: tabla_pers_sin_total.append({'Perfil': perfil, 'Descripción': fila.get('Descripción', ''), 'Cantidad': int(cantidad)})
            num_pers_total_int = int(datos_hecho.get('num_personal_capacitacion', 0))
            tabla_pers_data = tabla_pers_sin_total + [{'Perfil':'Total', 'Descripción':'', 'Cantidad': num_pers_total_int}]
            tabla_pers_subdoc_final = create_personal_table_subdoc(tpl_principal, ["Perfil", "Descripción", "Cantidad"], tabla_pers_data, ['Perfil', 'Descripción', 'Cantidad'], column_widths=(2,3,1), texto_posterior="Elaboración: Subdirección de Sanción y Gestión de Incentivos (SSAG) - DFAI.")

        # 4. Iterar sobre cada extremo para PREPARAR DATOS
        for j, extremo in enumerate(datos_hecho['extremos']):
            # a. Calcular CE unificado del extremo (Sin cambios)
            res_ce = _calcular_costo_evitado_extremo_inf005(datos_comunes, datos_hecho, extremo)
            if res_ce.get('error'): st.error(f"Error CE Extremo {j+1}: {res_ce['error']}"); continue
            aplicar_ce2_bi_extremo = res_ce['aplicar_ce2_a_bi'] # Flag específico del extremo

            # b. Calcular BI del extremo (Sin cambios)
            tipo_inc = extremo.get('tipo_extremo'); fecha_inc = extremo.get('fecha_incumplimiento'); fecha_ext = extremo.get('fecha_extemporanea')
            texto_bi = f"Extremo {j+1}: {extremo.get('tipo_monitoreo')}"
            datos_bi_base = {**datos_comunes, 'ce_soles': res_ce['ce_soles_para_bi'], 'ce_dolares': res_ce['ce_dolares_para_bi'], 'fecha_incumplimiento': fecha_inc, 'texto_del_hecho': texto_bi}
            res_bi_parcial = None
            if tipo_inc == "Remitió fuera de plazo":
                pre_bi = calcular_beneficio_ilicito(datos_bi_base);
                if pre_bi.get('error'): continue
                datos_bi_ext = {**datos_bi_base, 'fecha_cumplimiento_extemporaneo': fecha_ext, **pre_bi}; res_bi_parcial = calcular_beneficio_ilicito_extemporaneo(datos_bi_ext)
            else: res_bi_parcial = calcular_beneficio_ilicito(datos_bi_base)
            if not res_bi_parcial or res_bi_parcial.get('error'): st.warning(f"Error BI Extremo {j+1}: {res_bi_parcial.get('error', 'Error')}"); continue

            # c. Acumular totales (Sin cambios)
            bi_uit = res_bi_parcial.get('beneficio_ilicito_uit', 0.0); total_bi_uit += bi_uit
            anexos_ids.update(res_ce.get('ids_anexos', set()))
            resultados_app['totales']['ce1_total_soles'] += res_ce.get('ce1_soles', 0.0)
            resultados_app['totales']['ce2_total_soles_calculado'] += res_ce.get('ce2_soles_calculado', 0.0)
            resultados_app['totales']['ce_total_soles_para_bi'] += res_ce.get('ce_soles_para_bi')
            resultados_app['extremos'].append({'tipo': f"{extremo.get('tipo_monitoreo')} ({tipo_inc})", 'ce1_data': res_ce['ce1_data_raw'], 'ce2_data': res_ce['ce2_data_raw'], 'ce1_soles': res_ce['ce1_soles'], 'ce2_soles_calculado': res_ce['ce2_soles_calculado'], 'ce_soles_para_bi': res_ce['ce_soles_para_bi'], 'bi_data': res_bi_parcial.get('table_rows', []), 'bi_uit': bi_uit, 'aplicar_ce2_a_bi': aplicar_ce2_bi_extremo})
            lista_bi_resultados_completos.append(res_bi_parcial) # Guardar para multa final

            # d. Generar Anexo CE del extremo (Sin cambios)
            tpl_anx_loop = DocxTemplate(io.BytesIO(buffer_anexo.getvalue()))
            ce1_fmt_anx = []
            for item in res_ce['ce1_data_raw']:
                desc_orig = item.get('descripcion', '')
                texto_adicional = ""
                if "Profesional" in desc_orig: texto_adicional = "1/ "
                elif "Alquiler de laptop" in desc_orig: texto_adicional = "2/ "
                
                ce1_fmt_anx.append({
                    'descripcion': f"{texto_adicional}{desc_orig}",
                    'cantidad': format_decimal_dinamico(item.get('cantidad', 0)),
                    'horas': format_decimal_dinamico(item.get('horas', 0)),
                    'precio_soles': f"S/ {item.get('precio_soles', 0):,.3f}",
                    'factor_ajuste': f"{item.get('factor_ajuste', 0):,.3f}",
                    'monto_soles': f"S/ {item.get('monto_soles', 0):,.3f}",
                    'monto_dolares': f"US$ {item.get('monto_dolares', 0):,.3f}"
                })
            
            ce1_fmt_anx.append({
                'descripcion': 'Total',
                'monto_soles': f"S/ {res_ce['ce1_soles']:,.3f}",
                'monto_dolares': f"US$ {res_ce['ce1_dolares']:,.3f}"
            })
            
            tabla_ce1_anx = create_table_subdoc(
                tpl_anx_loop, 
                ["Descripción", "Cantidad", "Horas", "Precio asociado (S/)", "Factor de ajuste", "Monto (S/)", "Monto (US$)"],
                ce1_fmt_anx, 
                ['descripcion', 'cantidad', 'horas', 'precio_soles', 'factor_ajuste', 'monto_soles', 'monto_dolares']
            )
            tabla_ce2_anx = None
            if res_ce['ce2_data_raw']: 
                ce2_fmt_anx = [] # ce2_fmt_anx debe ser local
                for item in res_ce['ce2_data_raw']:
                    ce2_fmt_anx.append({
                        'descripcion': item.get('descripcion', ''),
                        'precio_dolares': f"US$ {item.get('precio_dolares', 0):,.3f}",
                        'precio_soles': f"S/ {item.get('precio_soles', 0):,.3f}",
                        'factor_ajuste': f"{item.get('factor_ajuste', 0):,.3f}",
                        'monto_soles': f"S/ {item.get('monto_soles', 0):,.3f}",
                        'monto_dolares': f"US$ {item.get('monto_dolares', 0):,.3f}"
                    })
                
                ce2_fmt_anx.append({
                    'descripcion': 'Total',
                    'monto_soles': f"S/ {res_ce['ce2_soles_calculado']:,.3f}",
                    'monto_dolares': f"US$ {res_ce['ce2_dolares_calculado']:,.3f}"
                })
                
                tabla_ce2_anx = create_table_subdoc(
                    tpl_anx_loop, 
                    ["Descripción", "Precio (US$)", "Precio (S/)", "Factor de ajuste", "Monto (S/)", "Monto (US$)"],
                    ce2_fmt_anx, 
                    ['descripcion', 'precio_dolares', 'precio_soles', 'factor_ajuste', 'monto_soles', 'monto_dolares']
                )
            
            fuentes_ce = res_ce.get('fuentes', {})
            contexto_anx = {**datos_comunes['context_data'], 'hecho': {'numero_imputado': num_hecho}, 'extremo': {'numeral': j+1, 'tipo': f"{extremo.get('tipo_monitoreo')}-{tipo_inc}", 'periodicidad': extremo.get('periodicidad',''), 'fecha_incumplimiento': format_date(fecha_inc, "d/MM/yyyy"), 'fecha_extemporanea': format_date(fecha_ext, "d/MM/yyyy") if fecha_ext else "N/A"}, 'tabla_ce1': tabla_ce1_anx, 'tabla_ce2': tabla_ce2_anx, 'aplicar_ce2_a_bi': aplicar_ce2_bi_extremo, # Fuentes de CE1 (Remisión)
                'fuente_salario_ce1': fuentes_ce.get('ce1', {}).get('fuente_salario', ''),
                'pdf_salario_ce1': fuentes_ce.get('ce1', {}).get('pdf_salario', ''),
                'sustento_prof_ce1': fuentes_ce.get('ce1', {}).get('sustento_item_profesional', ''),
                'fuente_coti_ce1': fuentes_ce.get('ce1', {}).get('fuente_coti', ''),
                
                # Fuentes de CE2 (Capacitación)
                'fuente_salario_ce2': fuentes_ce.get('ce2', {}).get('fuente_salario', ''),
                'pdf_salario_ce2': fuentes_ce.get('ce2', {}).get('pdf_salario', ''),
                'fuente_coti_ce2': fuentes_ce.get('ce2', {}).get('fuente_coti', ''),
                # Fuentes Comunes (IPC/TC de la fecha de incumplimiento)
                'fi_mes': fuentes_ce.get('fi_mes', ''),
                'fi_ipc': f"{fuentes_ce.get('fi_ipc', 0)}",
                'fi_tc': f"{fuentes_ce.get('fi_tc', 0)}",
                # --- FIN CORRECCIÓN 1 ---
            }
            tpl_anx_loop.render(contexto_anx, autoescape=True); buf_anx_final = io.BytesIO(); tpl_anx_loop.save(buf_anx_final); anexos_ce.append(buf_anx_final)

            # e. Generar tablas CE para el CUERPO (usando tpl_principal) (Sin cambios)
            tabla_ce1_cuerpo = create_table_subdoc(
                tpl_principal, 
                ["Descripción", "Cantidad", "Horas", "Precio asociado (S/)", "Factor de ajuste", "Monto (S/)", "Monto (US$)"],
                ce1_fmt_anx, # Reutiliza ce1_fmt_anx que ya está formateado
                ['descripcion', 'cantidad', 'horas', 'precio_soles', 'factor_ajuste', 'monto_soles', 'monto_dolares']
            )
            tabla_ce2_cuerpo = None
            # Re-utiliza la variable 'ce2_fmt_anx' que ya creamos y formateamos arriba
            if res_ce['ce2_data_raw']: 
                tabla_ce2_cuerpo = create_table_subdoc(
                    tpl_principal, 
                    ["Descripción", "Precio (US$)", "Precio (S/)", "Factor de ajuste", "Monto (S/)", "Monto (US$)"],
                    ce2_fmt_anx, 
                    ['descripcion', 'precio_dolares', 'precio_soles', 'factor_ajuste', 'monto_soles', 'monto_dolares']
                )

            # --- INICIO LÓGICA CORREGIDA: Preparar datos BI CON superíndices ---
            filas_bi_crudas_ext, fn_map_ext, fn_data_ext = res_bi_parcial.get('table_rows', []), res_bi_parcial.get('footnote_mapping', {}), res_bi_parcial.get('footnote_data', {})
            es_ext_iter = (tipo_inc == "Remitió fuera de plazo")

            # 1. Preparar lista de fuentes para el final de la tabla
            fn_list_ext = [f"({l}) {obtener_fuente_formateada(k, fn_data_ext, id_infraccion, es_ext_iter)}" for l, k in sorted(fn_map_ext.items())]
            fn_data_dict_ext = {'list': fn_list_ext, 'elaboration': 'Elaboración: Subdirección de Sanción y Gestión de Incentivos (SSAG) - DFAI.', 'style': 'FuenteTabla'}

            # 2. Preparar filas de datos CON superíndices (T y nota al pie)
            filas_bi_con_superindice = []
            for fila in filas_bi_crudas_ext:
                nueva_fila = fila.copy()
                ref_letra = nueva_fila.get('ref') # ej: (d)
                
                # Obtiene el texto base (ej: "...(1+COSm)")
                texto_base = str(nueva_fila.get('descripcion_texto', ''))
                
                # Obtiene el superíndice existente (ej: "T]")
                super_existente = str(nueva_fila.get('descripcion_superindice', ''))
                
                # Añade el nuevo superíndice (ej: "(d)")
                if ref_letra:
                    super_existente += f"({ref_letra})" # <-- Se vuelve "T](d)"

                nueva_fila['descripcion_texto'] = texto_base
                nueva_fila['descripcion_superindice'] = super_existente
                
                filas_bi_con_superindice.append(nueva_fila)
            # --- FIN LÓGICA CORREGIDA ---

            # Crear tabla BI del cuerpo usando los datos MODIFICADOS y pasando footnotes_data
            tabla_bi_cuerpo = create_main_table_subdoc(
                tpl_principal,
                ["Descripción", "Monto"],
                filas_bi_con_superindice, # <-- USA LAS FILAS CON SUPERÍNDICE
                keys=['descripcion_texto', 'monto'],
                footnotes_data=fn_data_dict_ext, # <-- Pasa las fuentes para ponerlas debajo
                column_widths=(5, 1)
            )

            # f. Añadir datos del extremo a la lista para el bucle (tabla_bi ahora incluye superíndices y footnotes)
            lista_extremos_plantilla_word.append({
                'loop_index': j + 1,
                'numeral': f"{num_hecho}.{j + 1}",
                'descripcion': f"Cálculo para el Extremo {j+1}: {extremo.get('tipo_monitoreo')} ({tipo_inc})",
                
                # --- INICIO ADICIÓN: Placeholders por Extremo ---
                'tipo_monitoreo': extremo.get('tipo_monitoreo', ''),
                'periodicidad': extremo.get('periodicidad', ''),
                # --- FIN ADICIÓN ---
                # --- INICIO ADICIÓN: Fechas Largas ---
                'fecha_incumplimiento_larga': (format_date(fecha_inc, "d 'de' MMMM 'de' yyyy", locale='es') if fecha_inc else "N/A").replace("septiembre", "setiembre").replace("Septiembre", "Setiembre"),
                'fecha_extemporanea_larga': (format_date(fecha_ext, "d 'de' MMMM 'de' yyyy", locale='es') if fecha_ext else "N/A").replace("septiembre", "setiembre").replace("Septiembre", "Setiembre"),
                # --- FIN ADICIÓN ---
                'tabla_ce1': tabla_ce1_cuerpo,
                'tabla_ce2': tabla_ce2_cuerpo, # Será None si no aplica
                'aplicar_ce2_a_bi': aplicar_ce2_bi_extremo, # Flag específico del extremo para el {% if %}
                'tabla_bi': tabla_bi_cuerpo, # <-- Este subdoc AHORA tiene superscripts en los datos y footnotes abajo
                'bi_uit_extremo': f"{bi_uit:,.3f} UIT",
                'ph_ipc_promedio_salario_ce1': fuentes_ce.get('ce1', {}).get('texto_ipc_costeo_salario', ''),
                # Ya no pasamos refs separadas
            })
        # --- FIN DEL BUCLE DE EXTREMOS ---

        # 5. Post-Cálculo: Multa Final (Sin cambios)
        # --- INICIO: Recuperar Factor F y Calcular Multa ---
        factor_f = datos_hecho.get('factor_f_calculado', 1.0)

        res_multa_final = calcular_multa({
            **datos_comunes, 
            'beneficio_ilicito': total_bi_uit,
            'factor_f': factor_f # <--- NUEVO
        })
        multa_final_uit = res_multa_final.get('multa_final_uit', 0.0)
        # --- FIN ---
    
        # --- INICIO: LÓGICA DE REDUCCIÓN Y TOPE (Múltiple) ---
        datos_hecho_completos = datos_comunes.get('datos_hecho_completos', {})
        aplica_reduccion_str = datos_hecho_completos.get('aplica_reduccion', 'No')
        porcentaje_str = datos_hecho_completos.get('porcentaje_reduccion', '0%')
        multa_con_reduccion_uit = multa_final_uit
        
        if aplica_reduccion_str == 'Sí':
            reduccion_factor = 0.5 if porcentaje_str == '50%' else 0.7
            multa_con_reduccion_uit = redondeo_excel(multa_final_uit * reduccion_factor, 3)

        infraccion_info = datos_comunes['df_tipificacion'][datos_comunes['df_tipificacion']['ID_Infraccion'] == id_infraccion]
        tope_multa_uit = float('inf')
        if not infraccion_info.empty and pd.notna(infraccion_info.iloc[0].get('Tope_Multa_Infraccion')):
            tope_multa_uit = float(infraccion_info.iloc[0]['Tope_Multa_Infraccion'])

        multa_final_del_hecho_uit = min(multa_con_reduccion_uit, tope_multa_uit)
        se_aplica_tope = multa_con_reduccion_uit > tope_multa_uit
        multa_reducida_uit = multa_con_reduccion_uit if aplica_reduccion_str == 'Sí' else multa_final_uit
        # --- FIN: LÓGICA DE REDUCCIÓN Y TOPE ---
        res_multa_final = calcular_multa({**datos_comunes, 'beneficio_ilicito': total_bi_uit})
        multa_final_uit = res_multa_final.get('multa_final_uit', 0.0)
        tabla_multa_final_subdoc = create_main_table_subdoc( tpl_principal, ["Componentes", "Monto"], res_multa_final.get('multa_data_raw', []), ['Componentes', 'Monto'], texto_posterior="Elaboración: Subdirección de Sanción y Gestión de Incentivos (SSAG) - DFAI.", estilo_texto_posterior='FuenteTabla', column_widths=(5, 1) )

        # --- ELIMINADOS FOOTNOTES CONSOLIDADOS ---

        # 6. Contexto Final y Renderizado (Sin footnotes_consolidadas)
        contexto_final = {
            **datos_comunes['context_data'], 'acronyms': datos_comunes['acronym_manager'],
            'hecho': {
                'numero_imputado': num_hecho,
                'descripcion': RichText(datos_hecho.get('texto_hecho', '')),
                'lista_extremos': lista_extremos_plantilla_word,
             },
            'numeral_hecho': f"IV.{num_hecho + 1}", # Suma 1 para empezar en IV.2
            'bi_uit_total': f"{total_bi_uit:,.3f} UIT",
            'mh_uit': f"{multa_final_uit:,.3f} UIT",
            # --- INICIO: PLACEHOLDERS DE REDUCCIÓN Y TOPE ---
            'aplica_reduccion': aplica_reduccion_str == 'Sí',
            'porcentaje_reduccion': porcentaje_str,
            'texto_reduccion': datos_hecho_completos.get('texto_reduccion', ''),
            'memo_num': datos_hecho_completos.get('memo_num', ''),
            'memo_fecha': format_date(datos_hecho_completos.get('memo_fecha'), "d 'de' MMMM 'de' yyyy", locale='es') if datos_hecho_completos.get('memo_fecha') else '',
            'escrito_num': datos_hecho_completos.get('escrito_num', ''),
            'escrito_fecha': format_date(datos_hecho_completos.get('escrito_fecha'), "d 'de' MMMM 'de' yyyy", locale='es') if datos_hecho_completos.get('escrito_fecha') else '',
            'multa_con_reduccion_uit': f"{multa_con_reduccion_uit:,.3f} UIT",
            'se_aplica_tope': se_aplica_tope,
            'tope_multa_uit': f"{tope_multa_uit:,.3f} UIT",
            # --- FIN: PLACEHOLDERS DE REDUCCIÓN Y TOPE ---
            'tabla_multa_final': tabla_multa_final_subdoc, # Tabla multa final
            # 'footnotes_consolidadas': footnotes_consolidadas_subdoc, # <-- ELIMINADO
            'tabla_detalle_personal': tabla_pers_subdoc_final, # Tabla personal (será None si no aplica)
            'usa_capacitacion': aplicar_capacitacion_general, # Flag para condicional de texto/tabla personal en plantilla
            'num_personal_total_texto': texto_con_numero(num_pers_total_int, 'f') if aplicar_capacitacion_general else '',
            'texto_explicacion_prorrateo': '', # Se genera en app.py
        }
        tpl_principal.render(contexto_final, autoescape=True);
        buf_final = io.BytesIO(); tpl_principal.save(buf_final)

        # 7. Preparar datos para App (Actualizar totales) (Sin cambios)
        resultados_app['totales'] = {**resultados_app['totales'], 'beneficio_ilicito_uit': total_bi_uit, 'multa_data_raw': res_multa_final.get('multa_data_raw', []), 'multa_final_uit': multa_final_uit, 'bi_data_raw': lista_bi_resultados_completos, 'tabla_personal_data': tabla_pers_data, # --- CORRECCIÓN 1: Pasar variables de reducción correctas ---
            'aplica_reduccion': aplica_reduccion_str,
            'porcentaje_reduccion': porcentaje_str,
            'multa_con_reduccion_uit': multa_con_reduccion_uit,
            'multa_reducida_uit': multa_reducida_uit,
            'multa_final_aplicada': multa_final_del_hecho_uit # <-- ESTA ES CLAVE
            # ----------------------------------------------------------
        }

        # 8. Devolver resultados (Sin footnotes_consolidadas)
        return {
            'doc_pre_compuesto': buf_final,
            'resultados_para_app': resultados_app,
            'texto_explicacion_prorrateo': '', # Se genera en app.py
            'tabla_detalle_personal': tabla_pers_subdoc_final if aplicar_capacitacion_general else None,
            'usa_capacitacion': aplicar_capacitacion_general,
            'es_extemporaneo': any(e.get('tipo_extremo') == 'Remitió fuera de plazo' for e in datos_hecho['extremos']),
            'anexos_ce_generados': anexos_ce,
            'ids_anexos': list(filter(None, anexos_ids)),
            'tabla_personal_data': tabla_pers_data if aplicar_capacitacion_general else []
        }
    except Exception as e:
        import traceback; traceback.print_exc()
        try: import streamlit as st; st.error(f"Error fatal en _procesar_hecho_multiple INF005: {e}")
        except ImportError: print(f"Error fatal en _procesar_hecho_multiple INF005: {e}")
        return {'error': f"Error fatal en _procesar_hecho_multiple INF005: {e}"}