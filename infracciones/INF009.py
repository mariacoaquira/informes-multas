import streamlit as st
import pandas as pd
import io
from babel.dates import format_date
from num2words import num2words
from docx import Document
from docxcompose.composer import Composer
from docxtpl import DocxTemplate, RichText
from datetime import date, timedelta
# (Se elimina 'holidays' ya que no se usa)
from jinja2 import Environment

# --- IMPORTACIONES DE MÓDULOS PROPIOS ---
from textos_manager import obtener_fuente_formateada
from funciones import (create_main_table_subdoc, create_table_subdoc,
                     texto_con_numero, create_footnotes_subdoc,
                     create_personal_table_subdoc, format_decimal_dinamico)
from sheets import (calcular_beneficio_ilicito, calcular_multa,
                    descargar_archivo_drive,
                    calcular_beneficio_ilicito_extemporaneo)
from funciones import create_main_table_subdoc, create_table_subdoc, texto_con_numero, create_footnotes_subdoc, format_decimal_dinamico, redondeo_excel

# ---------------------------------------------------------------------
# FUNCIÓN AUXILIAR DE FECHAS: REGISTRO ANUAL (INF009)
# ---------------------------------------------------------------------

def _calcular_fechas_registro_inf009(anio):
    """
    Calcula las fechas fijas para el registro anual.
    - Fecha Máxima: 31 de diciembre del año del registro.
    - Fecha Incumplimiento: 1 de enero del año siguiente.
    """
    if not anio or anio < 2000:
        return None, None

    try:
        # El plazo máximo es el último día del año del registro
        fecha_maxima_presentacion = date(anio, 12, 31)
        
        # El incumplimiento es el día siguiente, 1 de enero del próximo año
        fecha_incumplimiento = date(anio + 1, 1, 1)
            
        return fecha_maxima_presentacion, fecha_incumplimiento
    except ValueError:
        return None, None # En caso de un año inválido
    

# ---------------------------------------------------------------------
# FUNCIÓN AUXILIAR: CÁLCULO CE (SOLO CE1 - REGISTRO)
# ---------------------------------------------------------------------

def _calcular_costo_evitado_extremo_inf009(datos_comunes, extremo_data):
    """
    Calcula el CE (solo Costo de Administración de Registro) para un único extremo de INF009.
    """
    result = {
        'ce_data_raw': [], 'ce_soles': 0.0, 'ce_dolares': 0.0,
        'ids_anexos': set(),
        'fuentes': {'ce1': {}},
        'error': None
    }
    try:
        # --- 1. Datos del Extremo y Generales ---
        fecha_incumplimiento_extremo = extremo_data.get('fecha_incumplimiento')

        if not fecha_incumplimiento_extremo:
            raise ValueError("Falta la fecha de incumplimiento del extremo.")

        # --- 2. Unificar Fecha y Fuentes de Incumplimiento ---
        df_ind = datos_comunes.get('df_indices')
        if df_ind is None: raise ValueError("Faltan df_indices en datos_comunes.")
        
        fecha_final_dt = pd.to_datetime(fecha_incumplimiento_extremo, errors='coerce')
        if pd.isna(fecha_final_dt): raise ValueError(f"Fecha de incumplimiento inválida: {fecha_incumplimiento_extremo}")

        ipc_row = df_ind[df_ind['Indice_Mes'].dt.to_period('M') == fecha_final_dt.to_period('M')]
        if ipc_row.empty:
            raise ValueError(f"No se encontró IPC/TC para la fecha de incumplimiento {fecha_final_dt.strftime('%B %Y')}")
            
        ipc_inc, tc_inc = ipc_row.iloc[0]['IPC_Mensual'], ipc_row.iloc[0]['TC_Mensual']
        if pd.isna(ipc_inc) or pd.isna(tc_inc) or tc_inc == 0:
            raise ValueError("Valores IPC/TC inválidos para la fecha de incumplimiento.")

        result['fuentes']['fi_mes'] = format_date(fecha_final_dt, "MMMM 'de' yyyy", locale='es')
        result['fuentes']['fi_ipc'] = float(ipc_inc)
        result['fuentes']['fi_tc'] = float(tc_inc)
        
        # --- 3. Calcular CE1 (Administración de Registro - Lógica Interna Horas Fijas) ---
        fecha_calculo_ce1 = fecha_final_dt

        # --- INICIO Lógica interna para calcular CE1 ---
        def _calcular_ce1_interno(datos_comunes_ce1, fecha_final):
            res_int = {'items_calculados': [], 'error': None, 'fuentes': {}}
            try:
                df_items_inf = datos_comunes_ce1.get('df_items_infracciones')
                df_costos = datos_comunes_ce1.get('df_costos_items')
                df_coti = datos_comunes_ce1.get('df_coti_general')
                df_sal = datos_comunes_ce1.get('df_salarios_general')
                df_ind_ce1 = datos_comunes_ce1.get('df_indices')
                id_rubro_ce1 = datos_comunes_ce1.get('id_rubro_seleccionado')
                
                id_inf_ce1 = 'INF009' 

                if any(df is None for df in [df_items_inf, df_costos, df_coti, df_sal, df_ind_ce1]): raise ValueError("Faltan DataFrames CE1.")

                ipc_inc_ce1, tc_inc_ce1 = ipc_inc, tc_inc
                fecha_final_dt_ce1 = fecha_final

                fuentes_ce1 = {'placeholders_dinamicos': {}}; items_ce1 = []; sal_capturado = False
                receta_ce1 = df_items_inf[df_items_inf['ID_Infraccion'] == id_inf_ce1]
                if receta_ce1.empty: raise ValueError(f"No hay receta CE1 para {id_inf_ce1}")

                for _, item_receta in receta_ce1.iterrows():
                    id_item = item_receta['ID_Item_Infraccion']; desc_item = item_receta.get('Nombre_Item', 'N/A')
                    costos_posibles = df_costos[df_costos['ID_Item_Infraccion'] == id_item].copy();
                    if costos_posibles.empty: continue
                    tipo_item = item_receta.get('Tipo_Item'); df_candidatos = pd.DataFrame()
                    if tipo_item == 'Variable': 
                        id_rubro_str = str(id_rubro_ce1) if id_rubro_ce1 else ''; df_candidatos = costos_posibles[costos_posibles['ID_Rubro'].astype(str).str.contains(fr'\b{id_rubro_str}\b', regex=True, na=False)].copy() if id_rubro_str else pd.DataFrame();
                        if df_candidatos.empty: df_candidatos = costos_posibles[costos_posibles['ID_Rubro'].isin(['', 'nan', None])].copy()
                    elif tipo_item == 'Fijo': df_candidatos = costos_posibles.copy()
                    if df_candidatos.empty: continue
                    fechas_fuente = []
                    for _, cand in df_candidatos.iterrows():
                        id_gen = cand['ID_General']; fecha_f = pd.NaT
                        if pd.notna(id_gen):
                            if 'SAL' in id_gen: f = df_sal[df_sal['ID_Salario'] == id_gen]; fecha_f = pd.to_datetime(f"{int(f.iloc[0]['Costeo_Salario'])}-12-31", errors='coerce') if not f.empty else pd.NaT
                            elif 'COT' in id_gen: f = df_coti[df_coti['ID_Cotizacion'] == id_gen]; fecha_f = pd.to_datetime(f.iloc[0]['Fecha_Costeo'], errors='coerce') if not f.empty else pd.NaT
                        fechas_fuente.append(fecha_f)
                    df_candidatos['Fecha_Fuente'] = fechas_fuente; df_candidatos.dropna(subset=['Fecha_Fuente'], inplace=True)
                    if df_candidatos.empty: continue
                    fecha_naive = fecha_final_dt_ce1.tz_localize(None) if fecha_final_dt_ce1.tzinfo else fecha_final_dt_ce1
                    df_candidatos['Fecha_Fuente_Naive'] = df_candidatos['Fecha_Fuente'].apply(lambda x: x.tz_localize(None) if pd.notna(x) and x.tzinfo else x)
                    df_candidatos['Diferencia_Dias'] = (df_candidatos['Fecha_Fuente_Naive'] - fecha_naive).dt.days.abs()
                    costo_final = df_candidatos.loc[df_candidatos['Diferencia_Dias'].idxmin()]
                    id_gen = costo_final['ID_General']; fecha_f = costo_final['Fecha_Fuente']
                    ipc_cost, tc_cost = 0.0, 0.0
                    if pd.notna(id_gen) and 'SAL' in id_gen: idx_anio = df_ind_ce1[df_ind_ce1['Indice_Mes'].dt.year == fecha_f.year]; ipc_cost, tc_cost = (float(idx_anio['IPC_Mensual'].mean()), float(idx_anio['TC_Mensual'].mean())) if not idx_anio.empty else (0.0, 0.0)
                    elif pd.notna(id_gen) and 'COT' in id_gen: idx_row = df_ind_ce1[df_ind_ce1['Indice_Mes'].dt.to_period('M') == fecha_f.to_period('M')]; ipc_cost, tc_cost = (float(idx_row.iloc[0]['IPC_Mensual']), float(idx_row.iloc[0]['TC_Mensual'])) if not idx_row.empty else (0.0, 0.0)
                    if ipc_cost == 0 or pd.isna(ipc_cost): continue
                    if pd.notna(id_gen):
                        if 'COT' in id_gen: 
                            f_row=df_coti[df_coti['ID_Cotizacion']==id_gen]; sust=f_row.iloc[0].get('Fuente_Cotizacion') if not f_row.empty else None; fuentes_ce1.setdefault('fuente_coti',[]).append(sust) if sust else None
                        elif 'SAL' in id_gen and not sal_capturado: 
                            f_row=df_sal[df_sal['ID_Salario']==id_gen];
                            if not f_row.empty: 
                                fuentes_ce1['fuente_salario']=f_row.iloc[0].get('Fuente_Salario',''); fuentes_ce1['pdf_salario']=f_row.iloc[0].get('PDF_Salario',''); sal_capturado=True
                    if "Profesional" in desc_item: fuentes_ce1['sustento_item_profesional'] = costo_final.get('Sustento_Item', '')
                    try: 
                        key_ph=f"fuente_{desc_item.split()[0].lower().replace(':','')}"; fecha_ph=format_date(fecha_f, 'MMMM yyyy', locale='es').lower(); texto_ph=f"{desc_item}:\n{fecha_ph}, IPC={ipc_cost:,.3f}"; 
                        if 'placeholders_dinamicos' not in fuentes_ce1: fuentes_ce1['placeholders_dinamicos'] = {}
                        fuentes_ce1['placeholders_dinamicos'][key_ph]=texto_ph
                    except: pass
                    costo_orig=float(costo_final.get('Costo_Unitario_Item', 0.0)); moneda=costo_final.get('Moneda_Item')
                    if moneda!='S/' and (tc_cost==0 or pd.isna(tc_cost)): continue
                    precio_s=costo_orig if moneda=='S/' else costo_orig*tc_cost
                    factor=redondeo_excel(ipc_inc_ce1/ipc_cost, 3) if ipc_cost>0 else 0
                    cant=float(item_receta.get('Cantidad_Recursos', 1.0)); horas=float(item_receta.get('Cantidad_Horas', 1.0))
                    
                    monto_s = redondeo_excel(cant*horas*precio_s*factor, 3)
                    monto_d = redondeo_excel(monto_s/tc_inc_ce1 if tc_inc_ce1 > 0 else 0, 3)
                    
                    items_ce1.append({"descripcion": desc_item, "cantidad": cant, "horas": horas, "precio_soles": precio_s, "precio_dolares": round(precio_s / tc_inc_ce1 if tc_inc_ce1 > 0 else 0, 3), "factor_ajuste": factor, "monto_soles": monto_s, "monto_dolares": monto_d, "id_anexo": costo_final.get('ID_Anexo_Drive')})
                if 'fuente_coti' in fuentes_ce1: fuentes_ce1['fuente_coti'] = "\n".join(filter(None, set(fuentes_ce1['fuente_coti'])))
                res_int['items_calculados'] = items_ce1
                res_int['fuentes'] = fuentes_ce1
            except Exception as e_int: res_int['error'] = f"Error interno CE1: {e_int}"
            return res_int
        # --- FIN Lógica interna CE1 ---

        res_ce1 = _calcular_ce1_interno(datos_comunes, fecha_calculo_ce1)
        if res_ce1.get('error'):
            result['error'] = f"CE1: {res_ce1['error']}"
            return result
        
        result['ce_data_raw'] = res_ce1.get('items_calculados', [])
        result['ce_soles'] = sum(item.get('monto_soles', 0) for item in result['ce_data_raw'])
        result['ce_dolares'] = sum(item.get('monto_dolares', 0) for item in result['ce_data_raw'])
        result['ids_anexos'].update(item.get('id_anexo') for item in result['ce_data_raw'] if item.get('id_anexo'))
        result['fuentes']['ce1'] = res_ce1.get('fuentes', {})

        if not result['error']: result['error'] = None
        return result

    except Exception as e:
        import traceback; traceback.print_exc()
        result['error'] = f"Error crítico en _calcular_costo_evitado_extremo_inf009: {e}"
        return result


# ---------------------------------------------------------------------
# FUNCIÓN 2: RENDERIZAR INPUTS
# ---------------------------------------------------------------------

def renderizar_inputs_especificos(i, df_dias_no_laborables=None):
    st.markdown("##### Detalles del Incumplimiento: Registro Anual RRSS (INF009)")
    datos_hecho = st.session_state.imputaciones_data[i] 

    st.markdown("###### **Registros no administrados (Extremos)**")
    if 'extremos' not in datos_hecho: datos_hecho['extremos'] = [{}]
    if st.button("➕ Añadir Extremo", key=f"add_extremo_{i}"): datos_hecho['extremos'].append({}); st.rerun()

    for j, extremo in enumerate(datos_hecho['extremos']):
        # Hardcodeamos el tipo de incumplimiento ya que no hay extemporáneo
        extremo['tipo_extremo'] = "No contó/administró"
        
        with st.container(border=True):
            col_titulo, col_boton_eliminar = st.columns([0.85, 0.15])
            with col_titulo:
                st.markdown(f"**Extremo n.° {j + 1}**")
            with col_boton_eliminar:
                if st.button(f"🗑️", key=f"del_extremo_{i}_{j}"): datos_hecho['extremos'].pop(j); st.rerun()

            # 1. Año del registro
            extremo['anio'] = st.number_input("Año del Registro", min_value=2000, max_value=date.today().year,
                                             step=1, key=f"anio_{i}_{j}", value=extremo.get('anio', date.today().year - 1))
            
            # 2. Fechas de Supervisión (Mantenidas dentro del extremo)
            st.markdown("---")
            st.caption("Periodo de supervisión para este registro:")
            col_sup1, col_sup2 = st.columns(2)
            with col_sup1:
                extremo['fecha_supervision_inicio'] = st.date_input("Inicio supervisión", key=f"sup_ini_{i}_{j}",
                                                                  value=extremo.get('fecha_supervision_inicio'), format="DD/MM/YYYY")
            with col_sup2:
                extremo['fecha_supervision_fin'] = st.date_input("Fin supervisión", key=f"sup_fin_{i}_{j}",
                                                                value=extremo.get('fecha_supervision_fin'), format="DD/MM/YYYY")

            # 3. Lógica de Fechas de Incumplimiento (Fijas)
            if extremo.get('anio'):
                fecha_max, fecha_inc = _calcular_fechas_registro_inf009(extremo['anio'])
                extremo['fecha_maxima_presentacion'] = fecha_max
                extremo['fecha_incumplimiento'] = fecha_inc
                
                col_m1, col_m2 = st.columns(2)
                with col_m1: st.metric("Fecha Límite", "31/12/{}".format(extremo['anio']))
                with col_m2: st.metric("Fecha Incumplimiento", "01/01/{}".format(extremo['anio'] + 1))

            # Se elimina explícitamente cualquier referencia a fecha_extemporanea
            extremo['fecha_extemporanea'] = None

    return datos_hecho


# ---------------------------------------------------------------------
# FUNCIÓN 3: VALIDACIÓN DE INPUTS
# ---------------------------------------------------------------------
def validar_inputs(datos_hecho):
    if not datos_hecho.get('extremos'):
        st.warning("Debe añadir al menos un extremo.")
        return False
    
    for j, extremo in enumerate(datos_hecho.get('extremos', [])):
        if not all([
            extremo.get('anio'),
            extremo.get('fecha_supervision_inicio'),
            extremo.get('fecha_supervision_fin')
        ]):
            st.warning(f"Extremo {j+1}: Faltan datos (Año o Periodo de supervisión).")
            return False
    return True


# ---------------------------------------------------------------------
# FUNCIÓN 4: DESPACHADOR PRINCIPAL
# ---------------------------------------------------------------------
def procesar_infraccion(datos_comunes, datos_hecho):
    """
    Decide si procesar como hecho simple (1 extremo) o múltiple (>1 extremo).
    """
    num_extremos = len(datos_hecho.get('extremos', []))
    if num_extremos == 0: return {'error': 'No se ha registrado ningún extremo.'}
    elif num_extremos == 1: return _procesar_hecho_simple(datos_comunes, datos_hecho)
    else: return _procesar_hecho_multiple(datos_comunes, datos_hecho)


# ---------------------------------------------------------------------
# FUNCIÓN 5: PROCESAR HECHO SIMPLE
# ---------------------------------------------------------------------
def _procesar_hecho_simple(datos_comunes, datos_hecho):
    """
    Procesa un hecho INF009 con un único extremo.
    """
    try:
        jinja_env = Environment(trim_blocks=True, lstrip_blocks=True)
        
        # 1. Cargar plantillas BI y CE simples
        df_tipificacion, id_infraccion = datos_comunes['df_tipificacion'], 'INF009'
        filas_inf = df_tipificacion[df_tipificacion['ID_Infraccion'] == id_infraccion]
        if filas_inf.empty: return {'error': f"No se encontró ID '{id_infraccion}' en Tipificación."}
        fila_inf = filas_inf.iloc[0]
        id_tpl_bi, id_tpl_ce = fila_inf.get('ID_Plantilla_BI'), fila_inf.get('ID_Plantilla_CE')
        if not id_tpl_bi or not id_tpl_ce: return {'error': f'Faltan IDs plantilla simple (BI o CE) para {id_infraccion}.'}
        buf_bi, buf_ce = descargar_archivo_drive(id_tpl_bi), descargar_archivo_drive(id_tpl_ce)
        if not buf_bi or not buf_ce: return {'error': f'Fallo descarga plantilla simple para {id_infraccion}.'}
        doc_tpl_bi = DocxTemplate(buf_bi); tpl_anexo = DocxTemplate(buf_ce)

        # 2. Calcular CE
        extremo = datos_hecho['extremos'][0]
        res_ce = _calcular_costo_evitado_extremo_inf009(datos_comunes, extremo)
        if res_ce.get('error'): return {'error': f"Error CE: {res_ce['error']}"}

        # 3. Calcular BI y Multa
        tipo_inc, fecha_inc, fecha_ext = extremo.get('tipo_extremo'), extremo.get('fecha_incumplimiento'), extremo.get('fecha_extemporanea')
        texto_bi = f"{datos_hecho.get('texto_hecho', 'Hecho no especificado')}"
        datos_bi_base = {**datos_comunes, 'ce_soles': res_ce['ce_soles'], 'ce_dolares': res_ce['ce_dolares'], 'fecha_incumplimiento': fecha_inc, 'texto_del_hecho': texto_bi}
        res_bi = calcular_beneficio_ilicito_extemporaneo({**datos_bi_base, 'fecha_cumplimiento_extemporaneo': fecha_ext, **calcular_beneficio_ilicito(datos_bi_base)}) if tipo_inc == "Presentó fuera de plazo" else calcular_beneficio_ilicito(datos_bi_base)
        if not res_bi or res_bi.get('error'): return res_bi or {'error': 'Error BI.'}
        bi_uit = res_bi.get('beneficio_ilicito_uit', 0)
        res_multa = calcular_multa({**datos_comunes, 'beneficio_ilicito': bi_uit})
        multa_uit = res_multa.get('multa_final_uit', 0)

        # 4. Generar Tablas Cuerpo
        ce_fmt = []
        for i, item in enumerate(res_ce['ce_data_raw']):
            desc_orig = item.get('descripcion', '')
            texto_adicional = f"{i+1}/ " 
            ce_fmt.append({
                'descripcion': f"{desc_orig} {texto_adicional}",
                'cantidad': format_decimal_dinamico(item.get('cantidad', 0)),
                'horas': format_decimal_dinamico(item.get('horas', 0)),
                'precio_soles': f"S/ {item.get('precio_soles', 0):,.3f}",
                'factor_ajuste': f"{item.get('factor_ajuste', 0):,.3f}",
                'monto_soles': f"S/ {item.get('monto_soles', 0):,.3f}",
                'monto_dolares': f"US$ {item.get('monto_dolares', 0):,.3f}"
            })
        ce_fmt.append({'descripcion': 'Total', 'monto_soles': f"S/ {res_ce['ce_soles']:,.3f}", 'monto_dolares': f"US$ {res_ce['ce_dolares']:,.3f}"})
        tabla_ce = create_table_subdoc(doc_tpl_bi, ["Descripción", "Cantidad", "Horas", "Precio asociado (S/)", "Factor de ajuste 3/", "Monto (*) (S/)", "Monto (*) (US$) 4/"],
                                      ce_fmt, ['descripcion', 'cantidad', 'horas', 'precio_soles', 'factor_ajuste', 'monto_soles', 'monto_dolares'])

        filas_bi_crudas, fn_map, fn_data = res_bi.get('table_rows', []), res_bi.get('footnote_mapping', {}), res_bi.get('footnote_data', {})
        es_ext = (tipo_inc == "Presentó fuera de plazo")
        fn_list = [f"({l}) {obtener_fuente_formateada(k, fn_data, id_infraccion, es_ext)}" for l, k in sorted(fn_map.items())]
        fn_data_dict = {'list': fn_list, 'elaboration': 'Elaboración: Subdirección de Sanción y Gestión de Incentivos (SSAG) - DFAI.', 'style': 'FuenteTabla'}
        filas_bi_para_tabla = []
        for fila in filas_bi_crudas:
            nueva_fila = fila.copy()
            ref_letra = nueva_fila.get('ref')
            texto_base = str(nueva_fila.get('descripcion_texto', ''))
            super_existente = str(nueva_fila.get('descripcion_superindice', ''))
            if ref_letra: super_existente += f"({ref_letra})"
            nueva_fila['descripcion_texto'] = texto_base
            nueva_fila['descripcion_superindice'] = super_existente
            filas_bi_para_tabla.append(nueva_fila)
        tabla_bi = create_main_table_subdoc(doc_tpl_bi, ["Descripción", "Monto"], filas_bi_para_tabla, keys=['descripcion_texto', 'monto'], footnotes_data=fn_data_dict, column_widths=(5.5, 0.5))

        tabla_multa = create_main_table_subdoc(doc_tpl_bi, ["Componentes", "Monto"], res_multa.get('multa_data_raw', []), ['Componentes', 'Monto'], texto_posterior="Elaboración: Subdirección de Sanción y Gestión de Incentivos (SSAG) - DFAI.", estilo_texto_posterior='FuenteTabla', column_widths=(5.5, 0.5))

        # 5. Contexto y Renderizado Cuerpo
        fuentes_ce = res_ce.get('fuentes', {})

        datos_hecho_completos = datos_comunes.get('datos_hecho_completos', {})
        aplica_reduccion_str = datos_hecho_completos.get('aplica_reduccion', 'No')
        porcentaje_str = datos_hecho_completos.get('porcentaje_reduccion', '0%')
        multa_con_reduccion_uit = multa_uit
        if aplica_reduccion_str == 'Sí':
            reduccion_factor = 0.5 if porcentaje_str == '50%' else 0.7
            multa_con_reduccion_uit = redondeo_excel(multa_uit * reduccion_factor, 3)
        infraccion_info = datos_comunes['df_tipificacion'][datos_comunes['df_tipificacion']['ID_Infraccion'] == id_infraccion]
        tope_multa_uit = float('inf') 
        if not infraccion_info.empty and pd.notna(infraccion_info.iloc[0].get('Tope_Multa_Infraccion')):
            tope_multa_uit = float(infraccion_info.iloc[0]['Tope_Multa_Infraccion'])
        multa_final_del_hecho_uit = min(multa_con_reduccion_uit, tope_multa_uit)
        se_aplica_tope = multa_con_reduccion_uit > tope_multa_uit
        multa_reducida_uit = multa_con_reduccion_uit if aplica_reduccion_str == 'Sí' else multa_uit
        
        contexto_word = {
            **datos_comunes['context_data'],
            **fuentes_ce.get('ce1', {}).get('placeholders_dinamicos', {}),
            'acronyms': datos_comunes['acronym_manager'],
            'hecho': {'numero_imputado': datos_comunes['numero_hecho_actual'], 'descripcion': RichText(datos_hecho.get('texto_hecho', ''))},
            'numeral_hecho': f"IV.{datos_comunes['numero_hecho_actual'] + 1}",
            'anio_declaracion': f"{extremo.get('anio', 'N/A')}",
            
            # ... otros campos ...
            'fecha_supervision_inicio': format_date(extremo.get('fecha_supervision_inicio'), "d 'de' MMMM 'de' yyyy", locale='es'),
            'fecha_supervision_fin': format_date(extremo.get('fecha_supervision_fin'), "d 'de' MMMM 'de' yyyy", locale='es'),
            # ...
            'aplicar_capacitacion': False, 
            'label_ce_principal': "CE",     
            'tabla_ce1': tabla_ce,          
            'tabla_ce2': None,              
            'tabla_bi': tabla_bi,
            'tabla_multa': tabla_multa,
            'tabla_detalle_personal': None, 
            'num_personal_total_texto': '',
            'multa_original_uit': f"{multa_uit:,.3f} UIT",
            'mh_uit': f"{multa_final_del_hecho_uit:,.3f} UIT",
            'bi_uit': f"{bi_uit:,.3f} UIT",
            'fuente_cos': res_bi.get('fuente_cos', ''),
            'texto_explicacion_prorrateo': '',
            'aplica_reduccion': aplica_reduccion_str == 'Sí',
            'porcentaje_reduccion': porcentaje_str,
            'texto_reduccion': datos_hecho_completos.get('texto_reduccion', ''),
            'memo_num': datos_hecho_completos.get('memo_num', ''),
            'memo_fecha': format_date(datos_hecho_completos.get('memo_fecha'), "d 'de' MMMM 'de' yyyy", locale='es') if datos_hecho_completos.get('memo_fecha') else '',
            'escrito_num': datos_hecho_completos.get('escrito_num', ''),
            'escrito_fecha': format_date(datos_hecho_completos.get('escrito_fecha'), "d 'de' MMMM 'de' yyyy", locale='es') if datos_hecho_completos.get('escrito_fecha') else '',
            'multa_con_reduccion_uit': f"{multa_con_reduccion_uit:,.3f} UIT", 
            'se_aplica_tope': se_aplica_tope,
            'tope_multa_uit': f"{tope_multa_uit:,.3f} UIT",
        }

        doc_tpl_bi.render(contexto_word, autoescape=True, jinja_env=jinja_env)
        buf_final_hecho = io.BytesIO()
        doc_tpl_bi.save(buf_final_hecho)

        # 6. Generar Anexo CE
        anexos_ce = []
        tabla_ce_anx = create_table_subdoc(tpl_anexo, ["Descripción", "Cantidad", "Horas", "Precio asociado (S/)", "Factor de ajuste 3/", "Monto (*) (S/)", "Monto (*) (US$) 4/"],
                                          ce_fmt, ['descripcion', 'cantidad', 'horas', 'precio_soles', 'factor_ajuste', 'monto_soles', 'monto_dolares'])
        
        resumen_anexo = [{'desc': 'Costo Evitado Total', 'sol': f"S/ {res_ce['ce_soles']:,.3f}", 'dol': f"US$ {res_ce['ce_dolares']:,.3f}"}]
        tabla_resumen_anx = create_table_subdoc(tpl_anexo, ["Componente", "Monto (*) (S/)", "Monto (*) (US$)"], resumen_anexo, ['desc', 'sol', 'dol'])
        
        contexto_anx = {
            **contexto_word,
            'extremo': {
                 'tipo': f"Registro Anual {extremo.get('anio', 'N/A')}",
                 'periodicidad': f"Anual {extremo.get('anio', 'N/A')}",
                 'tipo_incumplimiento': tipo_inc,
                 'fecha_incumplimiento': format_date(fecha_inc, "d/MM/yyyy"),
                 'fecha_extemporanea': format_date(fecha_ext, "d/MM/yyyy") if fecha_ext else "N/A",
            },
            'tabla_ce1_anexo': tabla_ce_anx,
            'tabla_ce2_anexo': None,
            'tabla_resumen_anexo': tabla_resumen_anx,
            'fuente_salario_ce1': fuentes_ce.get('ce1', {}).get('fuente_salario', ''),
            'pdf_salario_ce1': fuentes_ce.get('ce1', {}).get('pdf_salario', ''),
            'sustento_prof_ce1': fuentes_ce.get('ce1', {}).get('sustento_item_profesional', ''),
            'fuente_coti_ce1': fuentes_ce.get('ce1', {}).get('fuente_coti', ''),
            'fi_mes': fuentes_ce.get('fi_mes', ''),
            'fi_ipc': fuentes_ce.get('fi_ipc', 0),
            'fi_tc': fuentes_ce.get('fi_tc', 0),
        }
        tpl_anexo.render(contexto_anx, autoescape=True, jinja_env=jinja_env)
        buf_anexo_final = io.BytesIO()
        tpl_anexo.save(buf_anexo_final)
        anexos_ce.append(buf_anexo_final)

        # 7. Devolver Resultados
        resultados_app = {
             'totales': {
                  'ce_data_raw': res_ce['ce_data_raw'], 
                  'ce_total_soles': res_ce['ce_soles'],
                  'ce_total_dolares': res_ce['ce_dolares'],
                  'beneficio_ilicito_uit': bi_uit,
                  'multa_final_uit': multa_uit, 
                  'bi_data_raw': res_bi.get('table_rows', []),
                  'multa_data_raw': res_multa.get('multa_data_raw', []),
                  'aplica_reduccion': aplica_reduccion_str,
                  'porcentaje_reduccion': porcentaje_str,
                  'multa_con_reduccion_uit': multa_con_reduccion_uit,
                  'multa_reducida_uit': multa_reducida_uit, # <-- Añadido
                  'multa_final_aplicada': multa_final_del_hecho_uit 
             }
        }
        return {
            'doc_pre_compuesto': buf_final_hecho,
            'resultados_para_app': resultados_app,
            'es_extemporaneo': es_ext,
            'usa_capacitacion': False,
            'anexos_ce_generados': anexos_ce,
            'ids_anexos': list(filter(None, res_ce.get('ids_anexos', set()))),
            'texto_explicacion_prorrateo': '',
            'tabla_detalle_personal': None,
            'tabla_personal_data': []
        }

    except Exception as e:
        import traceback; traceback.print_exc()
        st.error(f"Error _procesar_simple INF009: {e}")
        return {'error': f"Error _procesar_simple INF009: {e}"}


# ---------------------------------------------------------------------
# FUNCIÓN 6: PROCESAR HECHO MÚLTIPLE
# ---------------------------------------------------------------------
def _procesar_hecho_multiple(datos_comunes, datos_hecho):
    """
    Procesa INF009 con múltiples extremos.
    """
    try:
        jinja_env = Environment(trim_blocks=True, lstrip_blocks=True)
        
        # 1. Cargar Plantillas
        df_tipificacion, id_infraccion = datos_comunes['df_tipificacion'], 'INF009'
        filas_inf = df_tipificacion[df_tipificacion['ID_Infraccion'] == id_infraccion]
        if filas_inf.empty: return {'error': f"No se encontró ID '{id_infraccion}' en Tipificación."}
        fila_inf = filas_inf.iloc[0]
        id_tpl_principal = fila_inf.get('ID_Plantilla_BI_Extremo')
        id_tpl_anx = fila_inf.get('ID_Plantilla_CE_Extremo')
        if not id_tpl_principal or not id_tpl_anx:
             return {'error': f'Faltan IDs de plantilla (BI_Extremo o CE_Extremo) para {id_infraccion}.'}
        buffer_plantilla = descargar_archivo_drive(id_tpl_principal)
        buffer_anexo = descargar_archivo_drive(id_tpl_anx)
        if not buffer_plantilla or not buffer_anexo:
            return {'error': f'Fallo descarga plantilla BI Extremo {id_tpl_principal} o anexo {id_tpl_anx}.'}
        tpl_principal = DocxTemplate(buffer_plantilla)

        # 2. Inicializar acumuladores
        total_bi_uit = 0.0; lista_bi_resultados_completos = []; anexos_ids = set()
        num_hecho = datos_comunes['numero_hecho_actual']; anexos_ce = []; lista_extremos_plantilla_word = []
        resultados_app = {'extremos': [], 'totales': {'ce_total_soles': 0, 'ce_total_dolares': 0, 'ce_data_raw': []}}
        
        # 4. Iterar sobre cada extremo
        for j, extremo in enumerate(datos_hecho['extremos']):
            # a. Calcular CE
            res_ce = _calcular_costo_evitado_extremo_inf009(datos_comunes, extremo)
            if res_ce.get('error'): st.error(f"Error CE Extremo {j+1}: {res_ce['error']}"); continue

            # b. Calcular BI
            tipo_inc, fecha_inc, fecha_ext = extremo.get('tipo_extremo'), extremo.get('fecha_incumplimiento'), extremo.get('fecha_extemporanea')
            texto_bi = f"{datos_hecho.get('texto_hecho', 'Hecho no especificado')} - Extremo {j + 1}"
            datos_bi_base = {**datos_comunes, 'ce_soles': res_ce['ce_soles'], 'ce_dolares': res_ce['ce_dolares'], 'fecha_incumplimiento': fecha_inc, 'texto_del_hecho': texto_bi}
            res_bi_parcial = calcular_beneficio_ilicito_extemporaneo({**datos_bi_base, 'fecha_cumplimiento_extemporaneo': fecha_ext, **calcular_beneficio_ilicito(datos_bi_base)}) if tipo_inc == "Presentó fuera de plazo" else calcular_beneficio_ilicito(datos_bi_base)
            if not res_bi_parcial or res_bi_parcial.get('error'): st.warning(f"Error BI Extremo {j+1}: {res_bi_parcial.get('error', 'Error')}"); continue

            # c. Acumular totales
            bi_uit = res_bi_parcial.get('beneficio_ilicito_uit', 0.0); total_bi_uit += bi_uit
            anexos_ids.update(res_ce.get('ids_anexos', set()))
            resultados_app['totales']['ce_total_soles'] += res_ce.get('ce_soles', 0.0)
            resultados_app['totales']['ce_total_dolares'] += res_ce.get('ce_dolares', 0.0)
            resultados_app['totales']['ce_data_raw'].extend(res_ce.get('ce_data_raw', []))
            resultados_app['extremos'].append({ 
                'tipo': f"Registro Anual {extremo.get('anio')} ({tipo_inc})",
                'ce_data': res_ce['ce_data_raw'], 
                'ce_soles': res_ce['ce_soles'], 'ce_dolares': res_ce['ce_dolares'],
                'bi_data': res_bi_parcial.get('table_rows', []), 'bi_uit': bi_uit, 
            })
            lista_bi_resultados_completos.append(res_bi_parcial)

            # d. Generar Anexo CE del extremo
            tpl_anx_loop = DocxTemplate(io.BytesIO(buffer_anexo.getvalue()))
            
            ce_fmt_anx = []
            for i, item in enumerate(res_ce['ce_data_raw']):
                desc_orig = item.get('descripcion', '')
                texto_adicional = f"{i+1}/ "
                ce_fmt_anx.append({
                    'descripcion': f"{texto_adicional}{desc_orig}", 'cantidad': format_decimal_dinamico(item.get('cantidad', 0)), 'horas': format_decimal_dinamico(item.get('horas', 0)),
                    'precio_soles': f"S/ {item.get('precio_soles', 0):,.3f}", 'factor_ajuste': f"{item.get('factor_ajuste', 0):,.3f}",
                    'monto_soles': f"S/ {item.get('monto_soles', 0):,.3f}", 'monto_dolares': f"US$ {item.get('monto_dolares', 0):,.3f}"
                })
            ce_fmt_anx.append({'descripcion': 'Total', 'monto_soles': f"S/ {res_ce['ce_soles']:,.3f}", 'monto_dolares': f"US$ {res_ce['ce_dolares']:,.3f}"})
            
            tabla_ce_anx = create_table_subdoc(
                tpl_anx_loop, 
                ["Descripción", "Cantidad", "Horas", "Precio asociado (S/)", "Factor de ajuste", "Monto (S/)", "Monto (US$)"],
                ce_fmt_anx, 
                ['descripcion', 'cantidad', 'horas', 'precio_soles', 'factor_ajuste', 'monto_soles', 'monto_dolares']
            )

            resumen_anexo = [{'desc': 'Costo Evitado Total', 'sol': f"S/ {res_ce['ce_soles']:,.3f}", 'dol': f"US$ {res_ce['ce_dolares']:,.3f}"}]
            tabla_resumen_anx = create_table_subdoc(tpl_anx_loop, ["Componente", "Monto (S/)", "Monto (US$)"], resumen_anexo, ['desc', 'sol', 'dol'])

            fuentes_ce = res_ce.get('fuentes', {})
            contexto_anx = {
                **datos_comunes['context_data'],
                **(fuentes_ce.get('ce1', {}).get('placeholders_dinamicos', {})),
                'acronyms': datos_comunes['acronym_manager'],
                'hecho': {'numero_imputado': num_hecho},
                'extremo': {
                    'numeral': j+1,
                    'tipo': f"Registro Anual {extremo.get('anio', 'N/A')}",
                    'periodicidad': f"Anual {extremo.get('anio', 'N/A')}",
                    'tipo_incumplimiento': tipo_inc,
                    'fecha_incumplimiento': format_date(fecha_inc, "d/MM/yyyy"),
                    'fecha_extemporanea': format_date(fecha_ext, "d/MM/yyyy") if fecha_ext else "N/A",
                },
                'tabla_ce1_anexo': tabla_ce_anx,
                'tabla_ce2_anexo': None,
                'tabla_resumen_anexo': tabla_resumen_anx,
                'aplicar_capacitacion': False,
                'fuente_salario_ce1': fuentes_ce.get('ce1', {}).get('fuente_salario', ''),
                'pdf_salario_ce1': fuentes_ce.get('ce1', {}).get('pdf_salario', ''),
                'sustento_prof_ce1': fuentes_ce.get('ce1', {}).get('sustento_item_profesional', ''),
                'fuente_coti_ce1': fuentes_ce.get('ce1', {}).get('fuente_coti', ''),
                'fi_mes': fuentes_ce.get('fi_mes', ''),
                'fi_ipc': fuentes_ce.get('fi_ipc', 0),
                'fi_tc': fuentes_ce.get('fi_tc', 0),
            }
            tpl_anx_loop.render(contexto_anx, autoescape=True, jinja_env=jinja_env); 
            buf_anx = io.BytesIO(); tpl_anx_loop.save(buf_anx); anexos_ce.append(buf_anx)

            # e. Generar tablas CE para el CUERPO
            tabla_ce_cuerpo = create_table_subdoc(
                tpl_principal, 
                ["Descripción", "Cantidad", "Horas", "Precio asociado (S/)", "Factor de ajuste", "Monto (S/)", "Monto (US$)"],
                ce_fmt_anx,
                ['descripcion', 'cantidad', 'horas', 'precio_soles', 'factor_ajuste', 'monto_soles', 'monto_dolares']
            )
            
            filas_bi_crudas_ext, fn_map_ext, fn_data_ext = res_bi_parcial.get('table_rows', []), res_bi_parcial.get('footnote_mapping', {}), res_bi_parcial.get('footnote_data', {})
            es_ext_iter = (tipo_inc == "Presentó fuera de plazo")
            fn_list_ext = [f"({l}) {obtener_fuente_formateada(k, fn_data_ext, id_infraccion, es_ext_iter)}" for l, k in sorted(fn_map_ext.items())]
            fn_data_dict_ext = {'list': fn_list_ext, 'elaboration': 'Elaboración: Subdirección de Sanción y Gestión de Incentivos (SSAG) - DFAI.', 'style': 'FuenteTabla'}
            filas_bi_con_superindice = []
            for fila in filas_bi_crudas_ext:
                nueva_fila = fila.copy(); ref_letra = nueva_fila.get('ref')
                texto_base = str(nueva_fila.get('descripcion_texto', '')); super_existente = str(nueva_fila.get('descripcion_superindice', ''))
                if ref_letra: super_existente += f"({ref_letra})"
                nueva_fila['descripcion_texto'] = texto_base; nueva_fila['descripcion_superindice'] = super_existente
                filas_bi_con_superindice.append(nueva_fila)
            tabla_bi_cuerpo = create_main_table_subdoc(tpl_principal, ["Descripción", "Monto"], filas_bi_con_superindice,
                                                     keys=['descripcion_texto', 'monto'], footnotes_data=fn_data_dict_ext, column_widths=(5.5, 0.5))

            # f. Añadir datos del extremo a la lista para el bucle
            lista_extremos_plantilla_word.append({
                'loop_index': j + 1,
                'numeral': f"{num_hecho}.{j + 1}",
                'descripcion': f"Cálculo para el Extremo {j+1}: Registro Anual {extremo.get('anio', 'N/A')} ({tipo_inc})",
                'label_ce_principal': "CE",
                'tabla_ce1': tabla_ce_cuerpo,
                'tabla_ce2': None,
                'aplicar_capacitacion': False,
                'tabla_bi': tabla_bi_cuerpo,
                'bi_uit_extremo': f"{bi_uit:,.3f} UIT",
                'texto_razonabilidad': RichText(""),
            })
        # --- FIN DEL BUCLE DE EXTREMOS ---

        # 5. Post-Cálculo: Multa Final
        if not lista_bi_resultados_completos: return {'error': 'No se pudo calcular BI para ningún extremo.'}
        res_multa_final = calcular_multa({**datos_comunes, 'beneficio_ilicito': total_bi_uit})
        multa_final_uit = res_multa_final.get('multa_final_uit', 0.0)

        datos_hecho_completos = datos_comunes.get('datos_hecho_completos', {})
        aplica_reduccion_str = datos_hecho_completos.get('aplica_reduccion', 'No')
        porcentaje_str = datos_hecho_completos.get('porcentaje_reduccion', '0%')
        multa_con_reduccion_uit = multa_final_uit
        if aplica_reduccion_str == 'Sí':
            reduccion_factor = 0.5 if porcentaje_str == '50%' else 0.7
            multa_con_reduccion_uit = redondeo_excel(multa_final_uit * reduccion_factor, 3)
        infraccion_info = datos_comunes['df_tipificacion'][datos_comunes['df_tipificacion']['ID_Infraccion'] == id_infraccion]
        tope_multa_uit = float('inf')
        if not infraccion_info.empty and pd.notna(infraccion_info.iloc[0].get('Tope_Multa_Infraccion')):
            tope_multa_uit = float(infraccion_info.iloc[0]['Tope_Multa_Infraccion'])
        multa_final_del_hecho_uit = min(multa_con_reduccion_uit, tope_multa_uit)
        se_aplica_tope = multa_con_reduccion_uit > tope_multa_uit
        multa_reducida_uit = multa_con_reduccion_uit if aplica_reduccion_str == 'Sí' else multa_final_uit
        
        tabla_multa_final_subdoc = create_main_table_subdoc( tpl_principal, ["Componentes", "Monto"], res_multa_final.get('multa_data_raw', []), ['Componentes', 'Monto'], texto_posterior="Elaboración: Subdirección de Sanción y Gestión de Incentivos (SSAG) - DFAI.", estilo_texto_posterior='FuenteTabla', column_widths=(5.5, 0.5) )
        
        # 6. Contexto Final y Renderizado
        contexto_final = {
            **datos_comunes['context_data'], 'acronyms': datos_comunes['acronym_manager'],
            'hecho': {
                'numero_imputado': num_hecho,
                'descripcion': RichText(datos_hecho.get('texto_hecho', '')),
                'lista_extremos': lista_extremos_plantilla_word,
             },
            'numeral_hecho': f"IV.{num_hecho + 1}",
            
            # --- INICIO DE LA MODIFICACIÓN: Nuevos Placeholders ---
            'fecha_supervision_inicio': format_date(datos_hecho.get('fecha_supervision_inicio'), "d 'de' MMMM 'de' yyyy", locale='es') if datos_hecho.get('fecha_supervision_inicio') else "N/A",
            'fecha_supervision_fin': format_date(datos_hecho.get('fecha_supervision_fin'), "d 'de' MMMM 'de' yyyy", locale='es') if datos_hecho.get('fecha_supervision_fin') else "N/A",
            # --- FIN DE LA MODIFICACIÓN ---

            'bi_uit_total': f"{total_bi_uit:,.3f} UIT",
            'multa_original_uit': f"{multa_final_uit:,.3f} UIT",
            'mh_uit': f"{multa_final_del_hecho_uit:,.3f} UIT",
            'tabla_multa_final': tabla_multa_final_subdoc,
            'tabla_detalle_personal': None,
            'se_usa_capacitacion': False,
            'num_personal_total_texto': '',
            'texto_explicacion_prorrateo': '',
            
            'aplica_reduccion': aplica_reduccion_str == 'Sí',
            'porcentaje_reduccion': porcentaje_str,
            'texto_reduccion': datos_hecho_completos.get('texto_reduccion', ''),
            'memo_num': datos_hecho_completos.get('memo_num', ''),
            'memo_fecha': format_date(datos_hecho_completos.get('memo_fecha'), "d 'de' MMMM 'de' yyyy", locale='es') if datos_hecho_completos.get('memo_fecha') else '',
            'escrito_num': datos_hecho_completos.get('escrito_num', ''),
            'escrito_fecha': format_date(datos_hecho_completos.get('escrito_fecha'), "d 'de' MMMM 'de' yyyy", locale='es') if datos_hecho_completos.get('escrito_fecha') else '',
            'multa_con_reduccion_uit': f"{multa_con_reduccion_uit:,.3f} UIT",
            'se_aplica_tope': se_aplica_tope,
            'tope_multa_uit': f"{tope_multa_uit:,.3f} UIT",
        }

        tpl_principal.render(contexto_final, autoescape=True, jinja_env=jinja_env)
        buf_final = io.BytesIO(); tpl_principal.save(buf_final)

        # 7. Preparar datos para App
        resultados_app['totales'] = {
            **resultados_app['totales'], 
            'beneficio_ilicito_uit': total_bi_uit, 
            'multa_data_raw': res_multa_final.get('multa_data_raw', []), 
            'multa_final_uit': multa_final_uit, 
            'bi_data_raw': lista_bi_resultados_completos, 
            'tabla_personal_data': [],
            'aplica_reduccion': aplica_reduccion_str,
            'porcentaje_reduccion': porcentaje_str,
            'multa_con_reduccion_uit': multa_con_reduccion_uit, 
            'multa_reducida_uit': multa_reducida_uit,
            'multa_final_aplicada': multa_final_del_hecho_uit 
        }

        # 8. Devolver resultados
        return {
            'doc_pre_compuesto': buf_final,
            'resultados_para_app': resultados_app,
            'texto_explicacion_prorrateo': '',
            'tabla_detalle_personal': None,
            'usa_capacitacion': False,
            'es_extemporaneo': any(e.get('tipo_extremo') == 'Presentó fuera de plazo' for e in datos_hecho['extremos']),
            'anexos_ce_generados': anexos_ce,
            'ids_anexos': list(filter(None, anexos_ids)),
            'tabla_personal_data': [],
            'aplica_reduccion': aplica_reduccion_str,
            'porcentaje_reduccion': porcentaje_str,
            'multa_reducida_uit': multa_reducida_uit
        }
    except Exception as e:
        import traceback; traceback.print_exc(); return {'error': f"Error _procesar_multiple INF009: {e}"}
